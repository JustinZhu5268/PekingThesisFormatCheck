# -*- coding: utf-8 -*-
import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.oxml.ns import qn
import glob
import os

# Get latest fixed file
files = glob.glob('data/test_已修复_*.docx')
latest = max(files, key=os.path.getmtime)
print(f"Checking: {latest}")

# Check fixed doc
doc = Document(latest)

output = []

output.append("=" * 60)
output.append("FIXED: Paragraphs with sectPr")
output.append("=" * 60)

from docx.oxml import OxmlElement
for i, para in enumerate(doc.paragraphs):
    para_xml = para._element
    sectPr = para_xml.find(qn('w:sectPr'))
    if sectPr is not None:
        output.append(f"  段落 {i}: has sectPr - {para.text[:50]}...")

output.append(f"\nFIXED Total sections: {len(doc.sections)}")

# Check original
output.append("\n" + "=" * 60)
output.append("ORIGINAL: Paragraphs with sectPr")
output.append("=" * 60)

orig_file = latest.replace('.docx', '_原文件备份.docx')
if os.path.exists(orig_file):
    doc_orig = Document(orig_file)
    for i, para in enumerate(doc_orig.paragraphs):
        para_xml = para._element
        sectPr = para_xml.find(qn('w:sectPr'))
        if sectPr is not None:
            output.append(f"  段落 {i}: has sectPr - {para.text[:50]}...")
    output.append(f"\nORIGINAL Total sections: {len(doc_orig.sections)}")

with open('output_ch7.txt', 'w', encoding='utf-8') as f:
    f.write('\n'.join(output))
print("Done")
