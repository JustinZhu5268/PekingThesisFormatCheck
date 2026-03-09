# -*- coding: utf-8 -*-
import sys
sys.stdout.reconfigure(encoding='utf-8')

"""验证修复后的文档"""
from docx import Document

doc = Document('data/test_refs_fixed.docx')

print("=" * 60)
print("检查正文中的引用格式")
print("=" * 60)

# 查找正文中的引用
count = 0
for i in range(200, 300):
    para = doc.paragraphs[i]
    text = para.text
    if '（' in text and '）' in text:
        # 检查是否包含著者年份格式
        if '年' in text or any(c.isalpha() for c in text):
            print(f"段落 {i}: {text[:80]}")
            count += 1
            if count >= 10:
                break

print("\n" + "=" * 60)
print("检查参考文献标题后的段落")
print("=" * 60)

# 找到参考文献标题
ref_idx = None
for i, para in enumerate(doc.paragraphs):
    if para.text.strip() == "参考文献":
        ref_idx = i
        break

if ref_idx:
    print(f"参考文献标题在段落 {ref_idx}")
    print("\n前10条参考文献:")
    for i in range(ref_idx + 1, min(ref_idx + 11, len(doc.paragraphs))):
        para = doc.paragraphs[i]
        if para.text.strip():
            print(f"段落 {i}: {para.text[:100]}")
