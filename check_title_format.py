# -*- coding: utf-8 -*-
import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.shared import Pt

doc = Document('data/test_已修复_20260309_095405.docx')

# 找到参考文献、附录、致谢的位置
targets = {}
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if '参考文献' in text and len(text) < 10:
        targets['参考文献'] = i
    if text.startswith('附录'):
        targets['附录'] = i
    if '致谢' in text and len(text) < 10:
        targets['致谢'] = i

print('Found titles:')
for name, idx in targets.items():
    print(f'{name}: paragraph {idx}')

print('\n=== Title formatting comparison ===')
for name, idx in targets.items():
    para = doc.paragraphs[idx]
    print(f'\n{name} (para {idx}):')
    print(f'  text: "{para.text}"')
    print(f'  style: {para.style.name if para.style else "None"}')
    
    # 段落格式
    pf = para.paragraph_format
    print(f'  line_spacing: {pf.line_spacing}')
    print(f'  line_spacing_rule: {pf.line_spacing_rule}')
    print(f'  space_before: {pf.space_before.pt if pf.space_before else "None"} pt')
    print(f'  space_after: {pf.space_after.pt if pf.space_after else "None"} pt')
    print(f'  alignment: {pf.alignment}')
    
    # 检查runs
    for run_idx, run in enumerate(para.runs):
        print(f'  run {run_idx}: "{run.text}"')
        print(f'    font.size: {run.font.size.pt if run.font.size else "None"} pt')
        print(f'    font.bold: {run.font.bold}')
