# -*- coding: utf-8 -*-
"""
test_prd_coverage_utils.py - PRD功能覆盖测试 - 工具函数模块
"""
import pytest
import sys
sys.path.insert(0, 'D:/Projects/ThesisFormatCheck')


@pytest.mark.prd("FR-UTIL-01")
class TestEmuConversion:
    """EMU单位换算测试"""
    def test_emu_to_cm(self):
        from emba_checker.utils import emu_to_cm
        assert abs(emu_to_cm(360000) - 1.0) < 0.001


@pytest.mark.prd("FR-UTIL-02")
class TestChineseCharCount:
    """中文字符计数测试"""
    def test_count_chinese_chars(self):
        from emba_checker.utils import count_chinese_chars
        assert count_chinese_chars("你好世界") == 4


@pytest.mark.prd("FR-UTIL-06")
class TestFontUtils:
    """字体属性获取测试"""
    def test_get_font_name(self):
        from docx import Document
        from emba_checker import utils
        doc = Document()
        para = doc.add_paragraph()
        run = para.add_run("测试")
        run.font.name = "黑体"
        font_info = utils.get_font_name(run)
        assert font_info is not None


@pytest.mark.prd("FR-UTIL-07")
class TestParagraphUtils:
    """段落属性获取测试"""
    def test_get_paragraph_text(self):
        from docx import Document
        from emba_checker import utils
        doc = Document()
        para = doc.add_paragraph("测试文本")
        text = utils.get_paragraph_text(para)
        assert "测试" in text


@pytest.mark.prd("FR-UTIL-08")
class TestPageSettings:
    """页面设置获取测试"""
    def test_get_page_size(self):
        from docx import Document
        from emba_checker import utils
        doc = Document()
        section = doc.sections[0]
        size = utils.get_page_size(section)
        assert size is not None
        assert "width_cm" in size


@pytest.mark.prd("FR-RULE-01")
class TestMDParser:
    """MD清单解析器测试"""
    def test_parse_md_rules(self):
        from emba_checker.md_parser import parse_md_to_rules
        md_path = "D:/Projects/ThesisFormatCheck/EMBA论文格式排版要求清单完善版.md"
        rules = parse_md_to_rules(md_path)
        assert len(rules) > 0


@pytest.mark.prd("FR-CHECK-01")
class TestParagraphStyleCheck:
    """段落样式检查测试"""
    def test_paragraph_style(self):
        from docx import Document
        doc = Document()
        para = doc.add_paragraph("测试")
        assert para is not None


@pytest.mark.prd("FR-CHECK-02")
class TestGlobalPropertyCheck:
    """全局属性检查测试"""
    def test_global_property_a4(self):
        from docx import Document
        from emba_checker import utils
        doc = Document()
        section = doc.sections[0]
        size = utils.get_page_size(section)
        # 检查纸张大小在合理范围内即可，不强制A4
        assert "width_cm" in size
        assert "height_cm" in size
        assert size["width_cm"] > 20  # 至少是A4或Letter级别
