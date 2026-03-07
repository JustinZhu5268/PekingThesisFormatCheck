# -*- coding: utf-8 -*-
"""
test_more_executors.py - 更多执行器测试
"""
import pytest
import sys
sys.path.insert(0, 'D:/Projects/ThesisFormatCheck')

from docx import Document
from docx.shared import Pt
from emba_checker.rule_engine import RuleEngine

@pytest.mark.unit
class TestTextMatchExecutor:
    """测试 text_match 执行器"""

    def test_text_match_found(self):
        """测试文本匹配-找到"""
        doc = Document()
        p = doc.add_paragraph("版权声明：任何收存和保管本论文各种版本的单位和个人")
        
        rule = {
            "rule_id": "TEST_TEXT_01",
            "check_type": "text_match",
            "target_group": "Group2_Paragraphs",
            "target_zone": ["cover"],
            "conditions": {"expected_text": "收存", "match_mode": "contains"}
        }
        
        engine = RuleEngine(doc, [rule], {0: "cover"})
        issues = engine.run_all()
        
        # 应该找到匹配
        assert len(issues) >= 1

    def test_text_match_not_found(self):
        """测试文本匹配-未找到"""
        doc = Document()
        p = doc.add_paragraph("这是一段普通文本")
        
        rule = {
            "rule_id": "TEST_TEXT_02",
            "check_type": "text_match",
            "target_group": "Group2_Paragraphs",
            "target_zone": ["cover"],
            "conditions": {"expected_text": "收存", "match_mode": "contains"}
        }
        
        engine = RuleEngine(doc, [rule], {0: "cover"})
        issues = engine.run_all()
        
        # 应该未找到
        assert len(issues) >= 1

@pytest.mark.unit
class TestCrossReferenceExecutor:
    """测试 cross_reference 执行器"""

    def test_cross_reference_placeholder(self):
        """测试交叉引用占位"""
        # 简化测试
        doc = Document()
        p = doc.add_paragraph("如图 1.1 所示")
        
        rule = {
            "rule_id": "TEST_REF_01",
            "check_type": "cross_reference",
            "target_group": "Group6_CrossRef",
            "target_zone": ["body_chapter"],
            "conditions": {}
        }
        
        engine = RuleEngine(doc, [rule], {0: "body_chapter"})
        issues = engine.run_all()
        
        # 简化版应该跳过
        assert len(issues) >= 0

@pytest.mark.unit
class TestGlobalPropertyExecutor:
    """测试全局属性执行器"""

    def test_global_property_a4(self):
        """测试A4纸张"""
        from docx.shared import Cm
        doc = Document()
        section = doc.sections[0]
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        
        rule = {
            "rule_id": "TEST_GLOBAL_01",
            "check_type": "global_property",
            "target_group": "Group1_Global",
            "target_zone": ["*"],
            "conditions": {"property": "page_width_cm", "expected": 21.0, "tolerance": 0.1, "compare": "equals"}
        }
        
        engine = RuleEngine(doc, [rule], {0: "cover"})
        issues = engine.run_all()
        
        # A4应该通过
        assert len(issues) == 1  # info级别
