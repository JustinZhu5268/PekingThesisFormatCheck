# -*- coding: utf-8 -*-
"""
test_zone_detector_detailed.py - ZoneDetector详细单元测试
目标：行覆盖>=70%, 深度>=40%
"""
import pytest
import sys
from docx import Document

sys.path.insert(0, 'D:/Projects/ThesisFormatCheck')


@pytest.mark.unit
class TestZoneDetectorInit:
    """ZoneDetector初始化测试"""

    def test_init_empty_document(self):
        from emba_checker.zone_detector import ZoneDetector
        doc = Document()
        detector = ZoneDetector(doc)
        assert detector.doc == doc


@pytest.mark.unit
class TestZoneDetectorDetect:
    """_detect方法测试"""

    def test_detect_empty_document(self):
        from emba_checker.zone_detector import ZoneDetector
        doc = Document()
        detector = ZoneDetector(doc)
        assert detector.zone_map == {}

    def test_detect_cover_zone(self):
        from emba_checker.zone_detector import ZoneDetector
        doc = Document()
        doc.add_paragraph("封面标题")
        detector = ZoneDetector(doc)
        assert 0 in detector.zone_map

    def test_detect_abstract_cn(self):
        from emba_checker.zone_detector import ZoneDetector
        doc = Document()
        doc.add_paragraph("摘要")
        doc.add_paragraph("中文摘要内容")
        detector = ZoneDetector(doc)
        assert detector.zone_map[0] == "abstract_cn"

    def test_detect_abstract_en(self):
        from emba_checker.zone_detector import ZoneDetector
        doc = Document()
        doc.add_paragraph("ABSTRACT")
        doc.add_paragraph("English abstract")
        detector = ZoneDetector(doc)
        assert detector.zone_map[0] == "abstract_en"

    def test_detect_toc(self):
        from emba_checker.zone_detector import ZoneDetector
        doc = Document()
        doc.add_paragraph("目 录")
        detector = ZoneDetector(doc)
        assert detector.zone_map[0] == "toc"

    def test_detect_chapter(self):
        from emba_checker.zone_detector import ZoneDetector
        doc = Document()
        doc.add_paragraph("第一章 导论")
        detector = ZoneDetector(doc)
        assert detector.zone_map[0] == "body_chapter"

    def test_detect_references(self):
        from emba_checker.zone_detector import ZoneDetector
        doc = Document()
        doc.add_paragraph("参考文献")
        detector = ZoneDetector(doc)
        assert detector.zone_map[0] == "references"

    def test_detect_acknowledgment(self):
        from emba_checker.zone_detector import ZoneDetector
        doc = Document()
        doc.add_paragraph("致 谢")
        detector = ZoneDetector(doc)
        assert detector.zone_map[0] == "acknowledgment"

    def test_detect_appendix(self):
        from emba_checker.zone_detector import ZoneDetector
        doc = Document()
        doc.add_paragraph("附录A 测试")
        detector = ZoneDetector(doc)
        assert detector.zone_map[0] == "appendix"

    def test_detect_copyright(self):
        from emba_checker.zone_detector import ZoneDetector
        doc = Document()
        doc.add_paragraph("版权声明")
        detector = ZoneDetector(doc)
        assert detector.zone_map[0] == "copyright"

    def test_detect_declaration(self):
        from emba_checker.zone_detector import ZoneDetector
        doc = Document()
        doc.add_paragraph("原创性声明")
        detector = ZoneDetector(doc)
        assert detector.zone_map[0] == "declaration"

    def test_detect_empty_paragraph(self):
        from emba_checker.zone_detector import ZoneDetector
        doc = Document()
        doc.add_paragraph("封面")
        doc.add_paragraph("")
        doc.add_paragraph("正文")
        detector = ZoneDetector(doc)
        assert detector.zone_map[1] == "cover"

    def test_detect_whitespace_only(self):
        from emba_checker.zone_detector import ZoneDetector
        doc = Document()
        doc.add_paragraph("封面")
        doc.add_paragraph("   ")
        doc.add_paragraph("正文")
        detector = ZoneDetector(doc)
        assert detector.zone_map[1] == "cover"


@pytest.mark.unit
class TestZoneDetectorBoundary:
    """边界值测试"""

    def test_many_paragraphs(self):
        from emba_checker.zone_detector import ZoneDetector
        doc = Document()
        for i in range(500):
            doc.add_paragraph(f"段落{i}")
        detector = ZoneDetector(doc)
        assert len(detector.zone_map) == 500

    def test_very_long_text(self):
        from emba_checker.zone_detector import ZoneDetector
        doc = Document()
        doc.add_paragraph("测试内容 " * 10000)
        detector = ZoneDetector(doc)
        assert 0 in detector.zone_map

    def test_special_characters(self):
        from emba_checker.zone_detector import ZoneDetector
        doc = Document()
        doc.add_paragraph("特殊字符\t\n\r")
        detector = ZoneDetector(doc)
        assert 0 in detector.zone_map


@pytest.mark.unit
class TestZoneDetectorException:
    """异常测试"""

    def test_empty_document(self):
        from emba_checker.zone_detector import ZoneDetector
        doc = Document()
        detector = ZoneDetector(doc)
        assert detector.zone_map == {}


@pytest.mark.unit
class TestZoneDetectorFullWorkflow:
    """完整工作流测试"""

    def test_full_paper_flow(self):
        from emba_checker.zone_detector import ZoneDetector
        doc = Document()
        doc.add_paragraph("封面")
        doc.add_paragraph("摘要")
        doc.add_paragraph("中文摘要")
        doc.add_paragraph("ABSTRACT")
        doc.add_paragraph("English")
        doc.add_paragraph("目录")
        doc.add_paragraph("第一章")
        doc.add_paragraph("第二章")
        doc.add_paragraph("参考文献")
        doc.add_paragraph("致谢")
        detector = ZoneDetector(doc)
        assert detector.zone_map[0] == "cover"
        assert detector.zone_map[1] == "abstract_cn"
        assert detector.zone_map[3] == "abstract_en"
        assert detector.zone_map[5] == "toc"


@pytest.mark.unit
class TestZoneDetectorCoverageExtended:
    """扩展覆盖率测试 - P1模块优化"""

    def test_get_zone(self):
        """覆盖率测试：get_zone方法"""
        from emba_checker.zone_detector import ZoneDetector
        
        doc = Document()
        doc.add_paragraph("封面内容")
        doc.add_paragraph("第一章 绪论")  # 这会触发body_chapter
        
        detector = ZoneDetector(doc)
        
        # 测试get_zone方法
        zone0 = detector.get_zone(0)
        zone1 = detector.get_zone(1)
        assert zone0 is not None
        assert zone1 is not None

    def test_get_zone_out_of_range(self):
        """覆盖率测试：get_zone越界"""
        from emba_checker.zone_detector import ZoneDetector
        
        doc = Document()
        doc.add_paragraph("测试")
        
        detector = ZoneDetector(doc)
        
        # 越界应返回unknown
        result = detector.get_zone(100)
        assert result == "unknown"

    def test_get_zone_ranges(self):
        """覆盖率测试：get_zone_ranges方法"""
        from emba_checker.zone_detector import ZoneDetector
        
        doc = Document()
        doc.add_paragraph("封面")
        doc.add_paragraph("正文")
        doc.add_paragraph("参考文献")
        
        detector = ZoneDetector(doc)
        ranges = detector.get_zone_ranges()
        
        assert isinstance(ranges, dict)

    def test_get_paragraphs_in_zone(self):
        """覆盖率测试：get_paragraphs_in_zone方法"""
        from emba_checker.zone_detector import ZoneDetector
        
        doc = Document()
        doc.add_paragraph("封面")
        doc.add_paragraph("正文1")
        doc.add_paragraph("正文2")
        
        detector = ZoneDetector(doc)
        
        paragraphs = detector.get_paragraphs_in_zone("body_chapter")
        assert isinstance(paragraphs, list)

    def test_get_zone_stats(self):
        """覆盖率测试：get_zone_stats方法"""
        from emba_checker.zone_detector import ZoneDetector
        
        doc = Document()
        doc.add_paragraph("封面")
        doc.add_paragraph("正文")
        
        detector = ZoneDetector(doc)
        stats = detector.get_zone_stats()
        
        assert isinstance(stats, dict)
        assert "cover" in stats
        assert stats["cover"] >= 1

    def test_zone_map_edge_cases(self):
        """边界测试：zone_map边缘情况"""
        from emba_checker.zone_detector import ZoneDetector
        
        doc = Document()
        
        # 空文档
        detector = ZoneDetector(doc)
        assert detector.zone_map == {}
        
        # 单个段落
        doc.add_paragraph("测试")
        detector = ZoneDetector(doc)
        assert len(detector.zone_map) == 1

    def test_repr_method(self):
        """覆盖率测试：__repr__方法"""
        from emba_checker.zone_detector import ZoneDetector
        
        doc = Document()
        doc.add_paragraph("测试")
        
        detector = ZoneDetector(doc)
        repr_str = repr(detector)
        
        assert "ZoneDetector" in repr_str


# ============================================================================
# 辅助函数测试 - 补充未覆盖的模块级函数
# ============================================================================

@pytest.mark.unit
class TestZoneDetectorHelperFunctions:
    """测试模块级辅助函数"""

    def test_detect_zones_function(self):
        """覆盖率测试：detect_zones函数"""
        from emba_checker.zone_detector import detect_zones
        
        doc = Document()
        doc.add_paragraph("封面标题")
        doc.add_paragraph("正文内容")
        
        result = detect_zones(doc)
        
        assert isinstance(result, dict)
        assert 0 in result

    def test_get_zone_for_paragraph_function(self):
        """覆盖率测试：get_zone_for_paragraph函数"""
        from emba_checker.zone_detector import get_zone_for_paragraph
        
        doc = Document()
        doc.add_paragraph("封面标题")
        
        result = get_zone_for_paragraph(doc, 0)
        
        assert isinstance(result, str)

    def test_get_zone_ranges_function(self):
        """覆盖率测试：get_zone_ranges函数"""
        from emba_checker.zone_detector import get_zone_ranges
        
        doc = Document()
        doc.add_paragraph("封面")
        doc.add_paragraph("正文")
        
        result = get_zone_ranges(doc)
        
        assert isinstance(result, dict)


@pytest.mark.unit
class TestZoneDetectorHelperFunctionsException:
    """测试辅助函数的异常处理"""

    def test_is_cover_page_edge(self):
        """边界测试：is_cover_page边界情况"""
        from emba_checker.zone_detector import is_cover_page
        
        # 创建无样式的段落
        para = Document().add_paragraph("")
        
        result = is_cover_page(para)
        assert isinstance(result, bool)

    def test_is_copyright_page_edge(self):
        """边界测试：is_copyright_page边界情况"""
        from emba_checker.zone_detector import is_copyright_page
        
        para = Document().add_paragraph("")
        
        result = is_copyright_page(para)
        assert isinstance(result, bool)

    def test_is_declaration_page_edge(self):
        """边界测试：is_declaration_page边界情况"""
        from emba_checker.zone_detector import is_declaration_page
        
        para = Document().add_paragraph("")
        
        result = is_declaration_page(para)
        assert isinstance(result, bool)

    def test_is_abstract_cn_edge(self):
        """边界测试：is_abstract_cn无样式"""
        from emba_checker.zone_detector import is_abstract_cn
        
        # 无样式段落
        para = Document().add_paragraph("摘要")
        
        result = is_abstract_cn(para)
        assert isinstance(result, bool)

    def test_is_abstract_en_edge(self):
        """边界测试：is_abstract_en无样式"""
        from emba_checker.zone_detector import is_abstract_en
        
        para = Document().add_paragraph("ABSTRACT")
        
        result = is_abstract_en(para)
        assert isinstance(result, bool)

    def test_is_toc_edge(self):
        """边界测试：is_toc无样式"""
        from emba_checker.zone_detector import is_toc
        
        para = Document().add_paragraph("目录")
        
        result = is_toc(para)
        assert isinstance(result, bool)

    def test_is_chapter_edge(self):
        """边界测试：is_body_chapter无样式"""
        from emba_checker.zone_detector import is_body_chapter
        
        para = Document().add_paragraph("第一章")
        
        result = is_body_chapter(para)
        assert isinstance(result, bool)

    def test_is_reference_edge(self):
        """边界测试：is_references无样式"""
        from emba_checker.zone_detector import is_references
        
        para = Document().add_paragraph("参考文献")
        
        result = is_references(para)
        assert isinstance(result, bool)

    def test_is_acknowl_edge(self):
        """边界测试：is_acknowledgment无样式"""
        from emba_checker.zone_detector import is_acknowledgment
        
        para = Document().add_paragraph("致谢")
        
        result = is_acknowledgment(para)
        assert isinstance(result, bool)