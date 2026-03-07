# -*- coding: utf-8 -*-
"""
test_claude_engine_detailed.py - ClaudeEngine详细单元测试
目标：行覆盖>=70%, 深度>=40%
"""
import pytest
import sys
from unittest.mock import Mock, patch, MagicMock
from docx import Document

sys.path.insert(0, 'D:/Projects/ThesisFormatCheck')


@pytest.mark.unit
class TestClaudeEngineInit:
    """ClaudeEngine初始化测试"""

    def test_init_with_api_key(self):
        from emba_checker.claude_engine import ClaudeEngine
        doc = Document()
        engine = ClaudeEngine(doc, [], {}, api_key="test_key")
        assert engine.api_key == "test_key"

    def test_init_without_api_key(self):
        from emba_checker.claude_engine import ClaudeEngine
        doc = Document()
        engine = ClaudeEngine(doc, [], {}, api_key=None)
        assert engine.enabled == False

    def test_init_with_rules(self):
        from emba_checker.claude_engine import ClaudeEngine
        doc = Document()
        rules = [{"rule_id": "T1"}]
        engine = ClaudeEngine(doc, rules, {}, api_key=None)
        assert engine.rules == rules

    def test_init_with_zone_map(self):
        from emba_checker.claude_engine import ClaudeEngine
        doc = Document()
        zone_map = {0: "cover"}
        engine = ClaudeEngine(doc, [], zone_map, api_key=None)
        assert engine.zone_map == zone_map

    def test_init_with_progress_callback(self):
        from emba_checker.claude_engine import ClaudeEngine
        doc = Document()
        callback = Mock()
        engine = ClaudeEngine(doc, [], {}, api_key=None, progress_callback=callback)
        assert engine.progress_callback == callback


@pytest.mark.unit
class TestClaudeEngineRunAll:
    """run_all方法测试"""

    def test_run_all_disabled(self):
        from emba_checker.claude_engine import ClaudeEngine
        doc = Document()
        engine = ClaudeEngine(doc, [], {}, api_key=None)
        issues = engine.run_all()
        assert issues == []

    def test_run_all_empty_rules(self):
        from emba_checker.claude_engine import ClaudeEngine
        doc = Document()
        engine = ClaudeEngine(doc, [], {}, api_key=None)
        issues = engine.run_all()
        assert issues == []

    def test_run_all_disabled_rule(self):
        from emba_checker.claude_engine import ClaudeEngine
        doc = Document()
        rules = [{"rule_id": "T1", "enabled": False}]
        engine = ClaudeEngine(doc, rules, {}, api_key=None)
        issues = engine.run_all()
        assert isinstance(issues, list)


@pytest.mark.unit
class TestClaudeEngineProperties:
    """属性测试"""

    def test_default_model(self):
        from emba_checker.claude_engine import ClaudeEngine
        assert hasattr(ClaudeEngine, 'DEFAULT_MODEL')

    def test_timeout_constants(self):
        from emba_checker.claude_engine import ClaudeEngine
        assert hasattr(ClaudeEngine, 'SINGLE_CALL_TIMEOUT')
        assert hasattr(ClaudeEngine, 'TOTAL_TIMEOUT')

    def test_issues_property(self):
        from emba_checker.claude_engine import ClaudeEngine
        doc = Document()
        engine = ClaudeEngine(doc, [], {}, api_key=None)
        assert isinstance(engine.issues, list)

    def test_execution_log_property(self):
        from emba_checker.claude_engine import ClaudeEngine
        doc = Document()
        engine = ClaudeEngine(doc, [], {}, api_key=None)
        assert isinstance(engine.execution_log, list)

    def test_enabled_property(self):
        from emba_checker.claude_engine import ClaudeEngine
        doc = Document()
        engine = ClaudeEngine(doc, [], {}, api_key=None)
        assert engine.enabled == False

    def test_api_key_property(self):
        from emba_checker.claude_engine import ClaudeEngine
        doc = Document()
        engine = ClaudeEngine(doc, [], {}, api_key="test")
        assert engine.api_key == "test"


@pytest.mark.unit
class TestClaudeEngineException:
    """异常测试"""

    def test_client_init_error(self):
        from emba_checker.claude_engine import ClaudeEngine
        doc = Document()
        with patch.dict('sys.modules', {'anthropic': None}):
            engine = ClaudeEngine(doc, [], {}, api_key=None)
        assert engine.client is None

    def test_api_call_when_disabled(self):
        from emba_checker.claude_engine import ClaudeEngine
        doc = Document()
        engine = ClaudeEngine(doc, [], {}, api_key=None)
        issues = engine.run_all()
        assert isinstance(issues, list)
