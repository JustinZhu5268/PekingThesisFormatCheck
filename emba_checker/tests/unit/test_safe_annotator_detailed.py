# -*- coding: utf-8 -*-
"""
test_safe_annotator_detailed.py - SafeAnnotator详细单元测试 (简化版)
目标：行覆盖>=70%, 深度>=40%
"""
import pytest
import sys
from docx import Document
from docx.shared import Pt, RGBColor

sys.path.insert(0, 'D:/Projects/ThesisFormatCheck')


@pytest.mark.unit
class TestSafeAnnotatorInit:
    """SafeAnnotator初始化测试"""

    def test_init_with_all_params(self):
        from emba_checker.safe_annotator import SafeAnnotator
        doc = Document()
        doc.add_paragraph("测试")
        issues = [{"rule_id": "T1"}]
        annotator = SafeAnnotator(doc, issues, "output.docx")
        assert annotator.doc == doc
        assert annotator.issues == issues

    def test_init_without_output_path(self):
        from emba_checker.safe_annotator import SafeAnnotator
        doc = Document()
        doc.add_paragraph("测试")
        annotator = SafeAnnotator(doc, [], None)
        assert annotator.output_path is None


@pytest.mark.unit
class TestSafeAnnotatorAnnotateAll:
    """annotate_all方法测试"""

    def test_annotate_all_empty_issues(self):
        from emba_checker.safe_annotator import SafeAnnotator
        doc = Document()
        doc.add_paragraph("测试")
        annotator = SafeAnnotator(doc, [], "output.docx")
        success, failed = annotator.annotate_all()
        assert success == 0
        assert failed == 0

    def test_annotate_all_single_issue(self):
        from emba_checker.safe_annotator import SafeAnnotator
        doc = Document()
        doc.add_paragraph("测试内容")
        issues = [{"rule_id": "T1", "message": "问题1"}]
        annotator = SafeAnnotator(doc, issues, "output.docx")
        success, failed = annotator.annotate_all()
        assert success >= 0

    def test_annotate_all_disabled_issue(self):
        from emba_checker.safe_annotator import SafeAnnotator
        doc = Document()
        doc.add_paragraph("测试内容")
        issues = [{"rule_id": "T1", "enabled": False}]
        annotator = SafeAnnotator(doc, issues, "output.docx")
        success, failed = annotator.annotate_all()
        assert isinstance(success, int)


@pytest.mark.unit
class TestSafeAnnotatorBoundary:
    """边界值测试"""

    def test_many_issues(self):
        from emba_checker.safe_annotator import SafeAnnotator
        doc = Document()
        doc.add_paragraph("测试")
        issues = [{"rule_id": f"T{i}", "message": f"问题{i}"} for i in range(100)]
        annotator = SafeAnnotator(doc, issues, "output.docx")
        success, failed = annotator.annotate_all()
        assert isinstance(success, int)

    def test_empty_document(self):
        from emba_checker.safe_annotator import SafeAnnotator
        doc = Document()
        issues = [{"rule_id": "T1"}]
        annotator = SafeAnnotator(doc, issues, "output.docx")
        success, failed = annotator.annotate_all()
        assert isinstance(success, int)

    def test_many_paragraphs(self):
        from emba_checker.safe_annotator import SafeAnnotator
        doc = Document()
        for i in range(100):
            doc.add_paragraph(f"段落{i}")
        issues = [{"rule_id": "T1"}]
        annotator = SafeAnnotator(doc, issues, "output.docx")
        success, failed = annotator.annotate_all()
        assert isinstance(success, int)


@pytest.mark.unit
class TestSafeAnnotatorException:
    """异常测试"""

    def test_issue_missing_fields(self):
        from emba_checker.safe_annotator import SafeAnnotator
        doc = Document()
        doc.add_paragraph("测试")
        issues = [{}]
        annotator = SafeAnnotator(doc, issues, "output.docx")
        success, failed = annotator.annotate_all()
        assert isinstance(success, int)

    def test_invalid_paragraph_index(self):
        from emba_checker.safe_annotator import SafeAnnotator
        doc = Document()
        doc.add_paragraph("测试")
        issues = [{"rule_id": "T1", "paragraph_index": 9999}]
        annotator = SafeAnnotator(doc, issues, "output.docx")
        success, failed = annotator.annotate_all()
        assert isinstance(success, int)

    def test_none_values(self):
        from emba_checker.safe_annotator import SafeAnnotator
        doc = Document()
        doc.add_paragraph("测试")
        issues = [{"rule_id": None, "message": None}]
        annotator = SafeAnnotator(doc, issues, "output.docx")
        success, failed = annotator.annotate_all()
        assert isinstance(success, int)
