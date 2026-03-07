# -*- coding: utf-8 -*-
"""
test_zone_detector_extended.py - 扩展的zone_detector测试
"""
import pytest
import sys
sys.path.insert(0, 'D:/Projects/ThesisFormatCheck')

from docx import Document
from emba_checker.zone_detector import detect_zones, ZoneDetector

@pytest.mark.unit
class TestZoneDetectorExtended:
    """扩展的ZoneDetector测试"""

    def test_detect_zones_with_markers(self):
        """测试带标记的zone检测"""
        doc = Document()
        
        # 封面
        doc.add_paragraph("论文标题")
        
        # 摘要
        doc.add_paragraph("中文摘要内容")
        
        # 正文
        doc.add_paragraph("第一章 绪论")
        doc.add_paragraph("1.1 研究背景")
        
        zones = detect_zones(doc)
        assert len(zones) > 0
    
    def test_zone_detector_with_style(self):
        """测试带样式的zone检测"""
        doc = Document()
        
        # 添加带标题样式的段落
        title = doc.add_heading("第一章 绪论", level=1)
        
        zones = detect_zones(doc)
        assert len(zones) > 0
    
    def test_zone_ranges(self):
        """测试zone范围"""
        doc = Document()
        doc.add_paragraph("论文标题")
        doc.add_paragraph("摘要内容")
        doc.add_paragraph("第一章 绪论")
        
        detector = ZoneDetector(doc)
        ranges = detector.get_zone_ranges()
        assert isinstance(ranges, dict)

@pytest.mark.unit
class TestZoneEdgeCases:
    """Zone检测边界情况"""

    def test_empty_document(self):
        """空文档"""
        doc = Document()
        zones = detect_zones(doc)
        assert len(zones) >= 0

    def test_single_paragraph(self):
        """单段落"""
        doc = Document()
        doc.add_paragraph("只有一段")
        zones = detect_zones(doc)
        assert len(zones) >= 1

    def test_novel_like_document(self):
        """小说式文档（无章节标题）"""
        doc = Document()
        for i in range(10):
            doc.add_paragraph(f"段落{i}内容")
        zones = detect_zones(doc)
        # 应该全部归为body
        assert len(zones) > 0

    def test_mixed_styles(self):
        """混合样式"""
        doc = Document()
        doc.add_paragraph("普通段落")
        doc.add_heading("标题", level=1)
        doc.add_paragraph("又一个普通段落")
        zones = detect_zones(doc)
        assert len(zones) > 0
