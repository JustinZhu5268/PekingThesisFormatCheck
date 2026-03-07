# -*- coding: utf-8 -*-
"""
test_rule_engine_detailed.py - RuleEngine详细单元测试 (简化版)
目标：行覆盖>=80%, 深度>=60%
"""
import pytest
import sys
from docx import Document
from unittest.mock import Mock, patch, MagicMock

sys.path.insert(0, 'D:/Projects/ThesisFormatCheck')


@pytest.mark.unit
class TestRuleEngineInit:
    """RuleEngine初始化测试"""

    def test_init_with_all_params(self):
        from emba_checker.rule_engine import RuleEngine
        doc = Document()
        rules = [{"rule_id": "TEST_01"}]
        zone_map = {0: "cover"}
        engine = RuleEngine(doc, rules, zone_map, api_key="test_key")
        assert engine.doc == doc
        assert engine.rules == rules
        assert engine.zone_map == zone_map
        assert engine.api_key == "test_key"

    def test_init_without_api_key(self):
        from emba_checker.rule_engine import RuleEngine
        doc = Document()
        engine = RuleEngine(doc, [], {})
        assert engine.api_key is None

    def test_init_lazy_import(self):
        from emba_checker.rule_engine import RuleEngine
        RuleEngine.EXECUTOR_MAP = None
        doc = Document()
        engine = RuleEngine(doc, [], {})
        assert RuleEngine.EXECUTOR_MAP is not None


@pytest.mark.unit
class TestRuleEngineGetRules:
    """规则获取方法测试"""

    def test_get_rules_by_check_type(self):
        from emba_checker.rule_engine import RuleEngine
        doc = Document()
        rules = [
            {"rule_id": "T1", "check_type": "text_match"},
            {"rule_id": "T2", "check_type": "text_match"},
            {"rule_id": "T3", "check_type": "regex_match"},
        ]
        engine = RuleEngine(doc, rules, {})
        result = engine._get_rules_by_check_type("text_match")
        assert len(result) == 2

    def test_get_rules_by_check_type_not_found(self):
        from emba_checker.rule_engine import RuleEngine
        doc = Document()
        rules = [{"rule_id": "T1", "check_type": "text_match"}]
        engine = RuleEngine(doc, rules, {})
        result = engine._get_rules_by_check_type("nonexistent")
        assert result == []

    def test_get_rules_by_group(self):
        from emba_checker.rule_engine import RuleEngine
        doc = Document()
        rules = [
            {"rule_id": "T1", "target_group": "Group1_Global"},
            {"rule_id": "T2", "target_group": "Group1_Global"},
            {"rule_id": "T3", "target_group": "Group2_Paragraphs"},
        ]
        engine = RuleEngine(doc, rules, {})
        result = engine._get_rules_by_group("Group1_Global")
        assert len(result) == 2


@pytest.mark.unit
class TestRuleEngineRunAll:
    """run_all方法测试"""

    def test_run_all_empty_rules(self):
        from emba_checker.rule_engine import RuleEngine
        doc = Document()
        engine = RuleEngine(doc, [], {})
        issues = engine.run_all()
        assert issues == []

    def test_run_all_disabled_rule(self):
        from emba_checker.rule_engine import RuleEngine
        doc = Document()
        rules = [{"rule_id": "T1", "enabled": False}]
        engine = RuleEngine(doc, rules, {})
        issues = engine.run_all()
        assert issues == []

    def test_run_all_skip_check_type(self):
        from emba_checker.rule_engine import RuleEngine
        doc = Document()
        rules = [{"rule_id": "T1", "check_type": "skip"}]
        engine = RuleEngine(doc, rules, {})
        issues = engine.run_all()
        assert len(issues) == 1
        assert issues[0].result == "skip"

    def test_run_all_valid_rule(self):
        from emba_checker.rule_engine import RuleEngine
        doc = Document()
        rules = [{"rule_id": "T1", "check_type": "text_match", "enabled": True}]
        engine = RuleEngine(doc, rules, {})
        issues = engine.run_all()
        assert isinstance(issues, list)

    def test_run_all_missing_check_type(self):
        from emba_checker.rule_engine import RuleEngine
        doc = Document()
        rules = [{"rule_id": "T1"}]
        engine = RuleEngine(doc, rules, {})
        issues = engine.run_all()
        assert isinstance(issues, list)


@pytest.mark.unit
class TestRuleEngineRunByGroups:
    """run_by_groups方法测试"""

    def test_run_by_groups_empty(self):
        from emba_checker.rule_engine import RuleEngine
        doc = Document()
        engine = RuleEngine(doc, [], {})
        results = engine.run_by_groups()
        assert results == {}

    def test_run_by_groups_with_rules(self):
        from emba_checker.rule_engine import RuleEngine
        doc = Document()
        rules = [
            {"rule_id": "T1", "target_group": "Group1_Global", "check_type": "text_match"},
            {"rule_id": "T2", "target_group": "Group2_Paragraphs", "check_type": "text_match"},
        ]
        engine = RuleEngine(doc, rules, {})
        results = engine.run_by_groups()
        assert "Group1_Global" in results
        assert "Group2_Paragraphs" in results


@pytest.mark.unit
class TestRuleEngineRunSingleRule:
    """_run_single_rule方法测试"""

    def test_run_single_rule_skip(self):
        from emba_checker.rule_engine import RuleEngine
        doc = Document()
        rules = [{"rule_id": "T1", "check_type": "skip"}]
        engine = RuleEngine(doc, rules, {})
        issue = engine._run_single_rule(rules[0])
        assert issue.result == "skip"


@pytest.mark.unit
class TestRuleEngineExecutorMap:
    """执行器映射测试"""

    def test_executor_map_contains_all_types(self):
        from emba_checker.rule_engine import RuleEngine
        doc = Document()
        RuleEngine.EXECUTOR_MAP = None
        engine = RuleEngine(doc, [], {})
        expected_types = [
            "paragraph_style", "global_property", "regex_match",
            "count_check", "cross_reference", "text_match",
            "forbidden_words", "ai_semantic", "CheckIssue",
        ]
        for exec_type in expected_types:
            assert exec_type in RuleEngine.EXECUTOR_MAP


@pytest.mark.unit
class TestRuleEngineBoundary:
    """边界值测试"""

    def test_very_long_rules_list(self):
        from emba_checker.rule_engine import RuleEngine
        doc = Document()
        rules = [{"rule_id": f"T{i}", "check_type": "text_match"} for i in range(500)]
        engine = RuleEngine(doc, rules, {})
        issues = engine.run_all()
        assert isinstance(issues, list)

    def test_empty_document(self):
        from emba_checker.rule_engine import RuleEngine
        doc = Document()
        rules = [{"rule_id": "T1", "check_type": "text_match"}]
        engine = RuleEngine(doc, rules, {})
        issues = engine.run_all()
        assert isinstance(issues, list)

    def test_many_paragraphs(self):
        from emba_checker.rule_engine import RuleEngine
        doc = Document()
        for i in range(100):
            doc.add_paragraph(f"段落{i}")
        rules = [{"rule_id": "T1", "check_type": "text_match"}]
        engine = RuleEngine(doc, rules, {})
        issues = engine.run_all()
        assert isinstance(issues, list)

    def test_all_check_types(self):
        from emba_checker.rule_engine import RuleEngine
        doc = Document()
        check_types = [
            "paragraph_style", "global_property", "regex_match",
            "count_check", "cross_reference", "text_match",
            "forbidden_words", "ai_semantic"
        ]
        rules = [{"rule_id": f"T{i}", "check_type": ct} for i, ct in enumerate(check_types)]
        engine = RuleEngine(doc, rules, {})
        issues = engine.run_all()
        assert isinstance(issues, list)

    def test_mixed_enabled_rules(self):
        from emba_checker.rule_engine import RuleEngine
        doc = Document()
        rules = [
            {"rule_id": "T1", "check_type": "text_match", "enabled": True},
            {"rule_id": "T2", "check_type": "text_match", "enabled": False},
            {"rule_id": "T3", "check_type": "text_match"},
        ]
        engine = RuleEngine(doc, rules, {})
        issues = engine.run_all()
        assert isinstance(issues, list)


@pytest.mark.unit
class TestRuleEngineException:
    """异常测试"""

    def test_rule_with_none_enabled(self):
        from emba_checker.rule_engine import RuleEngine
        doc = Document()
        rules = [{"rule_id": "T1", "enabled": None}]
        engine = RuleEngine(doc, rules, {})
        issues = engine.run_all()
        assert isinstance(issues, list)

    def test_rule_with_none_check_type(self):
        from emba_checker.rule_engine import RuleEngine
        doc = Document()
        rules = [{"rule_id": "T1", "check_type": None}]
        engine = RuleEngine(doc, rules, {})
        issues = engine.run_all()
        assert isinstance(issues, list)

    def test_empty_rule(self):
        from emba_checker.rule_engine import RuleEngine
        doc = Document()
        engine = RuleEngine(doc, [{}], {})
        issues = engine.run_all()
        assert isinstance(issues, list)

    def test_executor_exception(self):
        from emba_checker.rule_engine import RuleEngine
        doc = Document()
        rules = [{"rule_id": "T1", "check_type": "text_match"}]
        engine = RuleEngine(doc, rules, {})
        with patch('emba_checker.executors.text_match_executor.TextMatchExecutor.check', 
                   side_effect=Exception("Test exception")):
            issues = engine.run_all()
        assert isinstance(issues, list)


@pytest.mark.unit
class TestRuleEngineIssues:
    """issues属性测试"""

    def test_issues_initially_empty(self):
        from emba_checker.rule_engine import RuleEngine
        doc = Document()
        engine = RuleEngine(doc, [], {})
        assert engine.issues == []

    def test_issues_accumulated(self):
        from emba_checker.rule_engine import RuleEngine
        doc = Document()
        rules = [{"rule_id": "T1", "check_type": "skip"}]
        engine = RuleEngine(doc, rules, {})
        engine.run_all()
        assert len(engine.issues) > 0


@pytest.mark.unit
class TestRuleEngineBoundaryExtended:
    """扩展边界测试 - P0模块优化"""

    def test_run_single_rule_boundary(self):
        """边界测试：_run_single_rule边界"""
        from emba_checker.rule_engine import RuleEngine
        
        doc = Document()
        doc.add_paragraph("测试")
        
        # 测试各种边界情况
        rules = [
            {"rule_id": "T1", "check_type": "text_match", "enabled": True},
            {"rule_id": "T2", "check_type": "text_match", "enabled": False},
            {"rule_id": "T3"},  # 缺少check_type
        ]
        
        engine = RuleEngine(doc, rules, {})
        
        # 直接测试_run_single_rule
        for rule in rules:
            result = engine._run_single_rule(rule)
            # 可能返回None、CheckIssue或list
            assert result is None or hasattr(result, 'rule_id') or isinstance(result, list)

    def test_get_rules_by_check_type_not_found(self):
        """边界测试：_get_rules_by_check_type找不到"""
        from emba_checker.rule_engine import RuleEngine
        
        doc = Document()
        rules = [{"rule_id": "T1", "check_type": "text_match"}]
        engine = RuleEngine(doc, rules, {})
        
        result = engine._get_rules_by_check_type("nonexistent_type")
        assert result == []

    def test_get_rules_by_group_not_found(self):
        """边界测试：_get_rules_by_group找不到"""
        from emba_checker.rule_engine import RuleEngine
        
        doc = Document()
        rules = [{"rule_id": "T1", "target_group": "Group1"}]
        engine = RuleEngine(doc, rules, {})
        
        result = engine._get_rules_by_group("NonExistentGroup")
        assert result == []


@pytest.mark.unit
class TestRuleEngineExceptionExtended:
    """扩展异常测试 - P0模块优化"""

    def test_executor_map_missing_type(self):
        """异常测试：executor_map中找不到类型"""
        from emba_checker.rule_engine import RuleEngine
        
        doc = Document()
        doc.add_paragraph("测试")
        
        # 传入不存在的check_type
        rules = [{"rule_id": "T1", "check_type": "completely_invalid_type_12345"}]
        engine = RuleEngine(doc, rules, {})
        issues = engine.run_all()
        
        assert isinstance(issues, list)

    def test_rule_with_missing_fields(self):
        """异常测试：规则缺少多个字段"""
        from emba_checker.rule_engine import RuleEngine
        
        doc = Document()
        doc.add_paragraph("测试")
        
        rules = [{}]  # 空规则
        engine = RuleEngine(doc, rules, {})
        issues = engine.run_all()
        
        assert isinstance(issues, list)

    def test_get_issue_summary_empty(self):
        """异常测试：get_issue_summary空结果"""
        from emba_checker.rule_engine import RuleEngine
        
        doc = Document()
        engine = RuleEngine(doc, [], {})
        
        summary = engine.get_issue_summary()
        assert isinstance(summary, dict)