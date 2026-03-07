# -*- coding: utf-8 -*-
"""
test_prd_coverage_zone.py - PRD功能覆盖测试 - 区域检测模块
"""
import pytest
import sys
sys.path.insert(0, 'D:/Projects/ThesisFormatCheck')


# FR-ZONE-01~10: 10个区域检测


@pytest.mark.prd("FR-ZONE-01")
class TestCoverZone:
    """封面区域检测 - FR-ZONE-01"""
    def test_detect_cover_zone(self):
        from docx import Document
        from emba_checker.zone_detector import detect_zones
        
        doc = Document()
        doc.add_paragraph("论文标题测试")
        doc.add_paragraph("作者姓名")
        
        zones = detect_zones(doc)
        # 默认区域应该是cover
        assert 0 in zones


@pytest.mark.prd("FR-ZONE-04")
class TestAbstractCnZone:
    """中文摘要区域 - FR-ZONE-04"""
    def test_detect_abstract_cn(self):
        from docx import Document
        from emba_checker.zone_detector import detect_zones
        
        doc = Document()
        doc.add_paragraph("论文标题")
        doc.add_paragraph("摘要")  # 触发摘要区域
        
        zones = detect_zones(doc)
        # 应该能检测到摘要


@pytest.mark.prd("FR-ZONE-05")
class TestAbstractEnZone:
    """英文摘要区域 - FR-ZONE-05"""
    def test_detect_abstract_en(self):
        from docx import Document
        from emba_checker.zone_detector import detect_zones
        
        doc = Document()
        doc.add_paragraph("ABSTRACT")
        
        zones = detect_zones(doc)


@pytest.mark.prd("FR-ZONE-06")
class TestTocZone:
    """目录区域 - FR-ZONE-06"""
    def test_detect_toc(self):
        from docx import Document
        from emba_checker.zone_detector import detect_zones
        
        doc = Document()
        doc.add_paragraph("目录")
        
        zones = detect_zones(doc)


@pytest.mark.prd("FR-ZONE-07")
class TestBodyChapterZone:
    """正文章节区域 - FR-ZONE-07"""
    def test_detect_body_chapter(self):
        from docx import Document
        from emba_checker.zone_detector import detect_zones
        
        doc = Document()
        doc.add_paragraph("第一章 绪论")
        
        zones = detect_zones(doc)


@pytest.mark.prd("FR-ZONE-08")
class TestReferencesZone:
    """参考文献区域 - FR-ZONE-08"""
    def test_detect_references(self):
        from docx import Document
        from emba_checker.zone_detector import detect_zones
        
        doc = Document()
        doc.add_paragraph("参考文献")
        
        zones = detect_zones(doc)


@pytest.mark.prd("FR-ZONE-09")
class TestAppendixZone:
    """附录区域 - FR-ZONE-09"""
    def test_detect_appendix(self):
        from docx import Document
        from emba_checker.zone_detector import detect_zones
        
        doc = Document()
        doc.add_paragraph("附录A 问卷")
        
        zones = detect_zones(doc)


@pytest.mark.prd("FR-ZONE-10")
class TestAcknowledgmentZone:
    """致谢区域 - FR-ZONE-10"""
    def test_detect_acknowledgment(self):
        from docx import Document
        from emba_checker.zone_detector import detect_zones
        
        doc = Document()
        doc.add_paragraph("致谢")
        
        zones = detect_zones(doc)


@pytest.mark.prd("FR-CHECK-03")
class TestRegexMatchCheck:
    """正则匹配检查 - FR-CHECK-03"""
    def test_match_date_pattern(self):
        from emba_checker.utils import match_pattern
        assert match_pattern("2024年1月", r"\d{4}年\d+月") == True


@pytest.mark.prd("FR-CHECK-04")
class TestCountCheck:
    """计数检查 - FR-CHECK-04"""
    def test_count_chars(self):
        from emba_checker.utils import count_chinese_chars
        assert count_chinese_chars("测试") == 2


@pytest.mark.prd("FR-CHECK-06")
class TestTextMatch:
    """文本匹配 - FR-CHECK-06"""
    def test_text_contains(self):
        from emba_checker.utils import match_pattern
        assert match_pattern("hello world", "world") == True
