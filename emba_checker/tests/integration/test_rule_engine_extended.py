# -*- coding: utf-8 -*-
"""
test_rule_engine_extended.py - RuleEngine扩展测试
包含边界测试和异常测试，用于提升测试覆盖率
"""
import pytest
import sys
from docx import Document
sys.path.insert(0, 'D:/Projects/ThesisFormatCheck')


@pytest.mark.integration
class TestRuleEngineBoundary:
    """RuleEngine边界测试"""

    def test_empty_rules_list(self):
        """边界测试：空规则列表"""
        from emba_checker.rule_engine import RuleEngine
        
        doc = Document()
        doc.add_paragraph("测试内容")
        
        engine = RuleEngine(doc, [], {})
        issues = engine.run_all()
        assert issues == []

    def test_single_rule(self):
        """边界测试：单条规则"""
        from emba_checker.rule_engine import RuleEngine
        
        doc = Document()
        doc.add_paragraph("测试内容")
        
        rules = [{
            "rule_id": "TEST_01",
            "check_type": "text_match",
            "enabled": True,
            "conditions": {"match_mode": "contains", "expected_text": "测试"}
        }]
        
        engine = RuleEngine(doc, rules, {})
        issues = engine.run_all()
        assert isinstance(issues, list)

    def test_many_rules(self):
        """边界测试：大量规则（100条）"""
        from emba_checker.rule_engine import RuleEngine
        
        doc = Document()
        doc.add_paragraph("测试内容")
        
        rules = [
            {
                "rule_id": f"TEST_{i:03d}",
                "check_type": "text_match",
                "enabled": True,
                "conditions": {"match_mode": "contains", "expected_text": "测试"}
            }
            for i in range(100)
        ]
        
        engine = RuleEngine(doc, rules, {})
        issues = engine.run_all()
        assert isinstance(issues, list)

    def test_empty_document(self):
        """边界测试：空文档"""
        from emba_checker.rule_engine import RuleEngine
        
        doc = Document()
        
        rules = [{
            "rule_id": "TEST_01",
            "check_type": "text_match",
            "enabled": True,
            "conditions": {"match_mode": "contains", "expected_text": "测试"}
        }]
        
        engine = RuleEngine(doc, rules, {})
        issues = engine.run_all()
        assert isinstance(issues, list)

    def test_many_paragraphs(self):
        """边界测试：多段落文档"""
        from emba_checker.rule_engine import RuleEngine
        
        doc = Document()
        for i in range(50):
            doc.add_paragraph(f"测试段落 {i}")
        
        rules = [{
            "rule_id": "TEST_01",
            "check_type": "text_match",
            "enabled": True,
            "conditions": {"match_mode": "contains", "expected_text": "测试"}
        }]
        
        engine = RuleEngine(doc, rules, {})
        issues = engine.run_all()
        assert isinstance(issues, list)

    def test_all_check_types(self):
        """边界测试：所有check_type"""
        from emba_checker.rule_engine import RuleEngine
        
        doc = Document()
        doc.add_paragraph("测试内容")
        
        rules = [
            {
                "rule_id": "TEST_01",
                "check_type": "paragraph_style",
                "enabled": True,
                "conditions": {}
            },
            {
                "rule_id": "TEST_02",
                "check_type": "global_property",
                "enabled": True,
                "conditions": {}
            },
            {
                "rule_id": "TEST_03",
                "check_type": "regex_match",
                "enabled": True,
                "conditions": {}
            },
            {
                "rule_id": "TEST_04",
                "check_type": "count_check",
                "enabled": True,
                "conditions": {}
            },
            {
                "rule_id": "TEST_05",
                "check_type": "cross_reference",
                "enabled": True,
                "conditions": {}
            },
            {
                "rule_id": "TEST_06",
                "check_type": "text_match",
                "enabled": True,
                "conditions": {}
            },
            {
                "rule_id": "TEST_07",
                "check_type": "forbidden_words",
                "enabled": True,
                "conditions": {}
            },
            {
                "rule_id": "TEST_08",
                "check_type": "ai_semantic",
                "enabled": True,
                "conditions": {}
            },
        ]
        
        engine = RuleEngine(doc, rules, {})
        issues = engine.run_all()
        assert isinstance(issues, list)


@pytest.mark.integration
class TestRuleEngineException:
    """RuleEngine异常测试"""

    def test_disabled_rule(self):
        """异常测试：禁用的规则"""
        from emba_checker.rule_engine import RuleEngine
        
        doc = Document()
        doc.add_paragraph("测试内容")
        
        rules = [{
            "rule_id": "TEST_01",
            "check_type": "text_match",
            "enabled": False,
            "conditions": {}
        }]
        
        engine = RuleEngine(doc, rules, {})
        issues = engine.run_all()
        assert isinstance(issues, list)

    def test_missing_check_type(self):
        """异常测试：缺少check_type"""
        from emba_checker.rule_engine import RuleEngine
        
        doc = Document()
        doc.add_paragraph("测试内容")
        
        rules = [{
            "rule_id": "TEST_01",
            "enabled": True,
            "conditions": {}
        }]
        
        engine = RuleEngine(doc, rules, {})
        issues = engine.run_all()
        assert isinstance(issues, list)

    def test_invalid_check_type(self):
        """异常测试：无效的check_type"""
        from emba_checker.rule_engine import RuleEngine
        
        doc = Document()
        doc.add_paragraph("测试内容")
        
        rules = [{
            "rule_id": "TEST_01",
            "check_type": "invalid_type",
            "enabled": True,
            "conditions": {}
        }]
        
        engine = RuleEngine(doc, rules, {})
        issues = engine.run_all()
        assert isinstance(issues, list)

    def test_missing_rule_id(self):
        """异常测试：缺少rule_id"""
        from emba_checker.rule_engine import RuleEngine
        
        doc = Document()
        doc.add_paragraph("测试内容")
        
        rules = [{
            "check_type": "text_match",
            "enabled": True,
            "conditions": {}
        }]
        
        engine = RuleEngine(doc, rules, {})
        issues = engine.run_all()
        assert isinstance(issues, list)

    def test_empty_rule(self):
        """异常测试：空规则"""
        from emba_checker.rule_engine import RuleEngine
        
        doc = Document()
        doc.add_paragraph("测试内容")
        
        engine = RuleEngine(doc, [{}], {})
        issues = engine.run_all()
        assert isinstance(issues, list)

    def test_get_rules_by_check_type_empty(self):
        """测试：按check_type获取规则（无匹配）"""
        from emba_checker.rule_engine import RuleEngine
        
        doc = Document()
        
        rules = [{
            "rule_id": "TEST_01",
            "check_type": "text_match"
        }]
        
        engine = RuleEngine(doc, rules, {})
        result = engine._get_rules_by_check_type("nonexistent")
        assert result == []

    def test_get_rules_by_check_type_match(self):
        """测试：按check_type获取规则（匹配）"""
        from emba_checker.rule_engine import RuleEngine
        
        doc = Document()
        
        rules = [
            {"rule_id": "TEST_01", "check_type": "text_match"},
            {"rule_id": "TEST_02", "check_type": "text_match"},
            {"rule_id": "TEST_03", "check_type": "regex_match"},
        ]
        
        engine = RuleEngine(doc, rules, {})
        result = engine._get_rules_by_check_type("text_match")
        assert len(result) == 2

    def test_get_rules_by_group_empty(self):
        """测试：按group获取规则（无匹配）"""
        from emba_checker.rule_engine import RuleEngine
        
        doc = Document()
        
        rules = [{
            "rule_id": "TEST_01",
            "target_group": "Group2_Paragraphs"
        }]
        
        engine = RuleEngine(doc, rules, {})
        result = engine._get_rules_by_group("nonexistent")
        assert result == []

    def test_get_rules_by_group_match(self):
        """测试：按group获取规则（匹配）"""
        from emba_checker.rule_engine import RuleEngine
        
        doc = Document()
        
        rules = [
            {"rule_id": "TEST_01", "target_group": "Group2_Paragraphs"},
            {"rule_id": "TEST_02", "target_group": "Group2_Paragraphs"},
            {"rule_id": "TEST_03", "target_group": "Group1_Global"},
        ]
        
        engine = RuleEngine(doc, rules, {})
        result = engine._get_rules_by_group("Group2_Paragraphs")
        assert len(result) == 2

    def test_lazy_import(self):
        """测试：延迟导入"""
        from emba_checker.rule_engine import RuleEngine
        
        doc = Document()
        engine = RuleEngine(doc, [], {})
        
        # 验证EXECUTOR_MAP已初始化
        assert RuleEngine.EXECUTOR_MAP is not None
        assert "paragraph_style" in RuleEngine.EXECUTOR_MAP
        assert "text_match" in RuleEngine.EXECUTOR_MAP
        assert "CheckIssue" in RuleEngine.EXECUTOR_MAP

    def test_executor_map_keys(self):
        """测试：执行器映射包含所有类型"""
        from emba_checker.rule_engine import RuleEngine
        
        doc = Document()
        engine = RuleEngine(doc, [], {})
        
        expected_types = [
            "paragraph_style",
            "global_property",
            "regex_match",
            "count_check",
            "cross_reference",
            "text_match",
            "forbidden_words",
            "ai_semantic",
        ]
        
        for check_type in expected_types:
            assert check_type in RuleEngine.EXECUTOR_MAP

    def test_api_key_parameter(self):
        """测试：API key参数"""
        from emba_checker.rule_engine import RuleEngine
        
        doc = Document()
        
        # 不带api_key
        engine1 = RuleEngine(doc, [], {})
        assert engine1.api_key is None
        
        # 带api_key
        engine2 = RuleEngine(doc, [], {}, api_key="test_key")
        assert engine2.api_key == "test_key"


@pytest.mark.integration
class TestRuleEngineMethods:
    """RuleEngine方法测试"""

    def test_run_all_returns_list(self):
        """测试：run_all返回列表"""
        from emba_checker.rule_engine import RuleEngine
        
        doc = Document()
        doc.add_paragraph("测试内容")
        
        engine = RuleEngine(doc, [], {})
        issues = engine.run_all()
        
        assert isinstance(issues, list)
        assert issues == []

    def test_issues_accumulated(self):
        """测试：issues累积"""
        from emba_checker.rule_engine import RuleEngine
        
        doc = Document()
        doc.add_paragraph("测试内容")
        
        rules = [{
            "rule_id": "TEST_01",
            "check_type": "text_match",
            "enabled": True,
            "conditions": {"match_mode": "contains", "expected_text": "测试"}
        }]
        
        engine = RuleEngine(doc, rules, {})
        engine.run_all()
        
        assert isinstance(engine.issues, list)

    def test_zone_map_parameter(self):
        """测试：zone_map参数"""
        from emba_checker.rule_engine import RuleEngine
        
        doc = Document()
        doc.add_paragraph("测试内容")
        
        zone_map = {0: "cover", 1: "body_chapter"}
        engine = RuleEngine(doc, [], zone_map)
        
        assert engine.zone_map == zone_map
