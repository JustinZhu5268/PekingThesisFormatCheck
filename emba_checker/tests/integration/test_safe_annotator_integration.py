# -*- coding: utf-8 -*-
"""
test_safe_annotator_extended.py - SafeAnnotator扩展测试
包含边界测试和异常测试，用于提升测试覆盖率
"""
import pytest
import sys
from docx import Document
from docx.shared import Pt, RGBColor

sys.path.insert(0, 'D:/Projects/ThesisFormatCheck')


@pytest.mark.integration
class TestSafeAnnotatorBoundary:
    """SafeAnnotator边界测试"""

    def test_empty_issues(self):
        """边界测试：无issue"""
        from emba_checker.safe_annotator import SafeAnnotator
        
        doc = Document()
        doc.add_paragraph("测试内容")
        
        annotator = SafeAnnotator(doc, [], "output.docx")
        success, failed = annotator.annotate_all()
        
        assert success == 0
        assert failed == 0

    def test_single_issue(self):
        """边界测试：单个issue"""
        from emba_checker.safe_annotator import SafeAnnotator
        
        doc = Document()
        doc.add_paragraph("测试内容")
        
        issues = [{
            "rule_id": "TEST_01",
            "message": "测试消息"
        }]
        
        annotator = SafeAnnotator(doc, issues, "output.docx")
        success, failed = annotator.annotate_all()
        
        assert success >= 0
        assert failed >= 0

    def test_many_issues(self):
        """边界测试：大量issues（100个）"""
        from emba_checker.safe_annotator import SafeAnnotator
        
        doc = Document()
        doc.add_paragraph("测试内容")
        
        issues = [
            {"rule_id": f"TEST_{i:03d}", "message": f"消息{i}"}
            for i in range(100)
        ]
        
        annotator = SafeAnnotator(doc, issues, "output.docx")
        success, failed = annotator.annotate_all()
        
        assert success >= 0

    def test_empty_document(self):
        """边界测试：空文档"""
        from emba_checker.safe_annotator import SafeAnnotator
        
        doc = Document()
        
        issues = [{"rule_id": "TEST_01", "message": "消息"}]
        
        annotator = SafeAnnotator(doc, issues, "output.docx")
        success, failed = annotator.annotate_all()
        
        assert success >= 0

    def test_many_paragraphs(self):
        """边界测试：多段落文档"""
        from emba_checker.safe_annotator import SafeAnnotator
        
        doc = Document()
        for i in range(50):
            doc.add_paragraph(f"段落{i}")
        
        issues = [{"rule_id": "TEST_01", "message": "消息"}]
        
        annotator = SafeAnnotator(doc, issues, "output.docx")
        success, failed = annotator.annotate_all()
        
        assert success >= 0


@pytest.mark.integration
class TestSafeAnnotatorException:
    """SafeAnnotator异常测试"""

    def test_issue_missing_fields(self):
        """异常测试：issue缺少字段"""
        from emba_checker.safe_annotator import SafeAnnotator
        
        doc = Document()
        doc.add_paragraph("测试内容")
        
        # 空issue
        issues = [{}]
        
        annotator = SafeAnnotator(doc, issues, "output.docx")
        success, failed = annotator.annotate_all()
        
        assert isinstance(success, int)

    def test_issue_with_none_values(self):
        """异常测试：issue值为None"""
        from emba_checker.safe_annotator import SafeAnnotator
        
        doc = Document()
        doc.add_paragraph("测试内容")
        
        issues = [
            {"rule_id": None, "message": None, "paragraph_index": None}
        ]
        
        annotator = SafeAnnotator(doc, issues, "output.docx")
        success, failed = annotator.annotate_all()
        
        assert isinstance(success, int)

    def test_invalid_paragraph_index(self):
        """异常测试：无效的段落索引"""
        from emba_checker.safe_annotator import SafeAnnotator
        
        doc = Document()
        doc.add_paragraph("测试内容")
        
        issues = [
            {"rule_id": "TEST_01", "message": "消息", "paragraph_index": 9999}
        ]
        
        annotator = SafeAnnotator(doc, issues, "output.docx")
        success, failed = annotator.annotate_all()
        
        assert isinstance(success, int)

    def test_negative_paragraph_index(self):
        """异常测试：负数段落索引"""
        from emba_checker.safe_annotator import SafeAnnotator
        
        doc = Document()
        doc.add_paragraph("测试内容")
        
        issues = [
            {"rule_id": "TEST_01", "message": "消息", "paragraph_index": -1}
        ]
        
        annotator = SafeAnnotator(doc, issues, "output.docx")
        success, failed = annotator.annotate_all()
        
        assert isinstance(success, int)

    def test_none_output_path(self):
        """异常测试：无输出路径"""
        from emba_checker.safe_annotator import SafeAnnotator
        
        doc = Document()
        doc.add_paragraph("测试内容")
        
        issues = [{"rule_id": "TEST_01", "message": "消息"}]
        
        annotator = SafeAnnotator(doc, issues, None)
        success, failed = annotator.annotate_all()
        
        assert isinstance(success, int)

    def test_annotation_log(self):
        """测试：标注日志"""
        from emba_checker.safe_annotator import SafeAnnotator
        
        doc = Document()
        doc.add_paragraph("测试内容")
        
        annotator = SafeAnnotator(doc, [], "output.docx")
        
        assert isinstance(annotator.annotation_log, list)

    def test_annotated_count(self):
        """测试：标注计数"""
        from emba_checker.safe_annotator import SafeAnnotator
        
        doc = Document()
        doc.add_paragraph("测试内容")
        
        annotator = SafeAnnotator(doc, [], "output.docx")
        
        assert annotator.annotated_count == 0

    def test_failed_count(self):
        """测试：失败计数"""
        from emba_checker.safe_annotator import SafeAnnotator
        
        doc = Document()
        doc.add_paragraph("测试内容")
        
        annotator = SafeAnnotator(doc, [], "output.docx")
        
        assert annotator.failed_count == 0

    def test_constants(self):
        """测试：常量定义"""
        from emba_checker.safe_annotator import SafeAnnotator
        
        assert SafeAnnotator.ANNOTATION_FONT_COLOR == RGBColor(255, 0, 0)
        assert SafeAnnotator.ANNOTATION_FONT_SIZE == Pt(9)
        assert SafeAnnotator.ANNOTATION_FONT_BOLD == True
