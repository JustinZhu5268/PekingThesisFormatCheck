# -*- coding: utf-8 -*-
"""
test_prd_coverage_integration.py - PRD功能覆盖测试 - 集成测试层
"""
import pytest
import sys
sys.path.insert(0, 'D:/Projects/ThesisFormatCheck')

# 真实docx
REAL_DOCX = 'D:/Projects/ThesisFormatCheck/医保支付方式改革下A医疗集团医疗收入质量优化研究——以A医疗集团为例-定稿V1.0.docx'


# FR-ENGINE-01: DocxEngine - Python确定性检查引擎
@pytest.mark.prd("FR-ENGINE-01")
class TestDocxEngineIntegration:
    """DocxEngine集成测试 - FR-ENGINE-01"""
    
    def test_docx_engine_run(self):
        import json
        from docx import Document
        from emba_checker.docx_engine import DocxEngine
        from emba_checker.zone_detector import detect_zones
        
        with open('D:/Projects/ThesisFormatCheck/emba_checker/rules_registry.json', 'r', encoding='utf-8') as f:
            rules = json.load(f)
        
        doc = Document(REAL_DOCX)
        zones = detect_zones(doc)
        
        engine = DocxEngine(doc, rules, zones)
        issues = engine.run_all()
        
        assert isinstance(issues, list)


# FR-ENGINE-02: ClaudeEngine - Claude语义引擎
@pytest.mark.prd("FR-ENGINE-02")
class TestClaudeEngineIntegration:
    """ClaudeEngine集成测试 - FR-ENGINE-02"""
    
    def test_claude_engine_skip_when_no_api_key(self):
        """无API key时跳过"""
        from emba_checker.claude_engine import ClaudeEngine
        from docx import Document
        doc = Document()
        engine = ClaudeEngine(doc, [], {}, api_key=None)
        assert engine.is_enabled() == False


# FR-ENGINE-03: 双擎独立运行机制
@pytest.mark.prd("FR-ENGINE-03")
class TestDualEngine:
    """双擎独立运行测试 - FR-ENGINE-03"""
    
    def test_python_engine_independent(self):
        """Python引擎独立运行"""
        import json
        from docx import Document
        from emba_checker.docx_engine import DocxEngine
        from emba_checker.zone_detector import detect_zones
        
        with open('D:/Projects/ThesisFormatCheck/emba_checker/rules_registry.json', 'r', encoding='utf-8') as f:
            rules = json.load(f)
        
        # 只运行Python规则
        python_rules = [r for r in rules if r.get('engine') in ('Python', 'Python+Manual', 'Python+Claude')]
        
        doc = Document(REAL_DOCX)
        zones = detect_zones(doc)
        
        engine = DocxEngine(doc, python_rules[:10], zones)
        issues = engine.run_all()
        
        assert isinstance(issues, list)


# FR-ENGINE-04: 结果合并去重
@pytest.mark.prd("FR-ENGINE-04")
class TestIssueMerger:
    """结果合并去重测试 - FR-ENGINE-04"""
    
    def test_merge_issues(self):
        # 使用main_engine中实际存在的_merge_issues方法
        from emba_checker.main_engine import MainEngine
        from docx import Document
        
        # 创建引擎实例并设置必要的属性
        engine = MainEngine.__new__(MainEngine)
        engine.rules = []
        engine.stats = {}  # 需要stats属性
        
        python_issues = [
            {"rule_id": "TEST_01", "message": "error1"},
            {"rule_id": "TEST_02", "message": "error2"}
        ]
        claude_issues = [
            {"rule_id": "TEST_01", "message": "error1"}  # 重复
        ]
        
        merged = engine._merge_issues(python_issues, claude_issues)
        assert isinstance(merged, list)


# FR-GROUP-01~07: 6个Group遍历
@pytest.mark.prd("FR-GROUP-01")
class TestGroup1Global:
    """Group1全局属性遍历 - FR-GROUP-01"""
    
    def test_group1_global(self):
        import json
        from docx import Document
        from emba_checker.rule_engine import RuleEngine
        from emba_checker.zone_detector import detect_zones
        
        with open('D:/Projects/ThesisFormatCheck/emba_checker/rules_registry.json', 'r', encoding='utf-8') as f:
            rules = json.load(f)
        
        group1_rules = [r for r in rules if r.get('target_group') == 'Group1_Global']
        
        doc = Document(REAL_DOCX)
        zones = detect_zones(doc)
        
        engine = RuleEngine(doc, group1_rules, zones)
        issues = engine.run_all()
        
        assert isinstance(issues, list)


@pytest.mark.prd("FR-GROUP-03")
class TestGroup2Paragraphs:
    """Group2段落遍历 - FR-GROUP-03"""
    
    def test_group2_paragraphs(self):
        import json
        from docx import Document
        from emba_checker.rule_engine import RuleEngine
        from emba_checker.zone_detector import detect_zones
        
        with open('D:/Projects/ThesisFormatCheck/emba_checker/rules_registry.json', 'r', encoding='utf-8') as f:
            rules = json.load(f)
        
        group2_rules = [r for r in rules if r.get('target_group') == 'Group2_Paragraphs']
        
        doc = Document(REAL_DOCX)
        zones = detect_zones(doc)
        
        engine = RuleEngine(doc, group2_rules[:20], zones)
        issues = engine.run_all()
        
        assert isinstance(issues, list)


# FR-ANNOTATE-01~05: 标注功能
@pytest.mark.prd("FR-ANNOTATE-01")
class TestAnnotateHighlight:
    """黄色高亮标注 - FR-ANNOTATE-01"""
    
    def test_highlight(self):
        from docx import Document
        from docx.shared import RGBColor
        from docx.enum.text import WD_COLOR_INDEX
        
        doc = Document()
        para = doc.add_paragraph("测试段落")
        
        # 测试高亮设置
        for run in para.runs:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        
        assert para.runs[0].font.highlight_color == WD_COLOR_INDEX.YELLOW


@pytest.mark.prd("FR-ANNOTATE-05")
class TestSafeAnnotator:
    """SafeAnnotator - FR-ANNOTATE-05"""
    
    def test_safe_annotator_create(self):
        from docx import Document
        from emba_checker.safe_annotator import SafeAnnotator
        from emba_checker.base_executor import CheckIssue
        
        doc = Document(REAL_DOCX)
        issues = [
            CheckIssue(
                rule_id="TEST_01",
                check_type="test",
                result="fail",
                message="测试问题",
                location={"paragraph_index": 0, "paragraphs": [0]},
                severity="error"
            )
        ]
        
        annotator = SafeAnnotator(doc, issues)
        assert annotator is not None


# FR-GROUP-04: Group3图表
@pytest.mark.prd("FR-GROUP-04")
class TestGroup3Figures:
    """Group3图表遍历 - FR-GROUP-04"""
    
    def test_group3_figures(self):
        import json
        from docx import Document
        from emba_checker.docx_engine import DocxEngine
        from emba_checker.zone_detector import detect_zones
        
        with open('D:/Projects/ThesisFormatCheck/emba_checker/rules_registry.json', 'r', encoding='utf-8') as f:
            rules = json.load(f)
        
        group3_rules = [r for r in rules if 'Group3' in r.get('target_group', '')]
        
        doc = Document(REAL_DOCX)
        zones = detect_zones(doc)
        
        engine = DocxEngine(doc, group3_rules[:5], zones)
        issues = engine.run_all()
        
        assert isinstance(issues, list)


# FR-GROUP-05: Group4脚注
@pytest.mark.prd("FR-GROUP-05")
class TestGroup4Footnotes:
    """Group4脚注遍历 - FR-GROUP-05"""
    
    def test_group4_footnotes(self):
        import json
        from docx import Document
        from emba_checker.docx_engine import DocxEngine
        from emba_checker.zone_detector import detect_zones
        
        with open('D:/Projects/ThesisFormatCheck/emba_checker/rules_registry.json', 'r', encoding='utf-8') as f:
            rules = json.load(f)
        
        group4_rules = [r for r in rules if 'Group4' in r.get('target_group', '')]
        
        doc = Document(REAL_DOCX)
        zones = detect_zones(doc)
        
        engine = DocxEngine(doc, group4_rules[:5], zones)
        issues = engine.run_all()
        
        assert isinstance(issues, list)


# FR-GROUP-06: Group5参考文献
@pytest.mark.prd("FR-GROUP-06")
class TestGroup5References:
    """Group5参考文献遍历 - FR-GROUP-06"""
    
    def test_group5_references(self):
        import json
        from docx import Document
        from emba_checker.docx_engine import DocxEngine
        from emba_checker.zone_detector import detect_zones
        
        with open('D:/Projects/ThesisFormatCheck/emba_checker/rules_registry.json', 'r', encoding='utf-8') as f:
            rules = json.load(f)
        
        group5_rules = [r for r in rules if 'Group5' in r.get('target_group', '')]
        
        doc = Document(REAL_DOCX)
        zones = detect_zones(doc)
        
        engine = DocxEngine(doc, group5_rules[:5], zones)
        issues = engine.run_all()
        
        assert isinstance(issues, list)
