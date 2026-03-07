# -*- coding: utf-8 -*-
"""
test_real_docx_integration.py - 使用真实docx文件的集成测试
"""
import pytest
import sys
import os
sys.path.insert(0, 'D:/Projects/ThesisFormatCheck')

# 真实docx文件路径
REAL_DOCX = 'D:/Projects/ThesisFormatCheck/data/医保支付方式改革下A医疗集团医疗收入质量优化研究——以A医疗集团为例-定稿V1.0.docx'

@pytest.mark.integration
class TestRealDocxIntegration:
    """使用真实docx文件的集成测试"""

    def test_load_real_docx(self):
        """测试加载真实docx文件"""
        from docx import Document
        doc = Document(REAL_DOCX)
        assert doc is not None
        assert len(doc.paragraphs) > 0

    def test_zone_detection_real_docx(self):
        """测试真实docx的区域检测"""
        from docx import Document
        from emba_checker.zone_detector import detect_zones
        
        doc = Document(REAL_DOCX)
        zones = detect_zones(doc)
        assert zones is not None
        assert len(zones) > 0

    def test_rule_engine_real_docx(self):
        """测试真实docx的规则引擎"""
        import json
        from docx import Document
        from emba_checker.rule_engine import RuleEngine
        from emba_checker.zone_detector import detect_zones
        
        # 加载规则
        with open('D:/Projects/ThesisFormatCheck/emba_checker/rules_registry.json', 'r', encoding='utf-8') as f:
            rules = json.load(f)
        
        doc = Document(REAL_DOCX)
        zones = detect_zones(doc)
        
        # 只测试前10条规则
        engine = RuleEngine(doc, rules[:10], zones)
        issues = engine.run_all()
        # 不崩溃即可
        assert isinstance(issues, list)

    def test_docx_engine_real_docx(self):
        """测试真实docx的DocxEngine"""
        import json
        from docx import Document
        from emba_checker.docx_engine import DocxEngine
        from emba_checker.zone_detector import detect_zones
        
        with open('D:/Projects/ThesisFormatCheck/emba_checker/rules_registry.json', 'r', encoding='utf-8') as f:
            rules = json.load(f)
        
        doc = Document(REAL_DOCX)
        zones = detect_zones(doc)
        
        engine = DocxEngine(doc, rules, zones)
        issues = engine.run_all()
        assert isinstance(issues, list)

    def test_safe_annotator_real_docx(self):
        """测试真实docx的SafeAnnotator"""
        from docx import Document
        from emba_checker.safe_annotator import SafeAnnotator
        from emba_checker.base_executor import CheckIssue
        
        doc = Document(REAL_DOCX)
        
        # 创建模拟的issue
        issues = [
            CheckIssue(
                rule_id="TEST_01",
                check_type="test",
                result="fail",
                message="测试问题",
                location={"paragraph_index": 0, "paragraphs": [0]},
                severity="error"
            )
        ]
        
        annotator = SafeAnnotator(doc, issues)
        result = annotator.annotate_all()
        assert isinstance(result, tuple)

    def test_main_engine_real_docx(self):
        """测试真实docx的MainEngine"""
        from emba_checker.main_engine import MainEngine
        
        engine = MainEngine(REAL_DOCX)
        assert engine is not None
        assert engine.doc is not None

    def test_main_engine_run(self):
        """测试MainEngine运行"""
        from emba_checker.main_engine import MainEngine
        
        engine = MainEngine(REAL_DOCX)
        issues = engine.run(enable_claude=False)
        assert isinstance(issues, list)

    def test_utils_with_real_docx(self):
        """测试真实docx的工具函数"""
        from docx import Document
        from emba_checker import utils
        
        doc = Document(REAL_DOCX)
        
        # 测试段落工具
        if doc.paragraphs:
            para = doc.paragraphs[0]
            text = utils.get_paragraph_text(para)
            assert text is not None
            
            is_empty = utils.is_paragraph_empty(para)
            assert isinstance(is_empty, bool)
            
            runs = list(utils.iter_paragraph_runs(para))
            assert isinstance(runs, list)

    def test_page_settings_real_docx(self):
        """测试真实docx的页面设置"""
        from docx import Document
        from emba_checker import utils
        
        doc = Document(REAL_DOCX)
        section = doc.sections[0]
        
        page_size = utils.get_page_size(section)
        assert page_size is not None
        
        margins = utils.get_page_margins(section)
        assert margins is not None
