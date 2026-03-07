# -*- coding: utf-8 -*-
"""
test_zone_detector_extended.py - ZoneDetector扩展测试
包含边界测试和异常测试，用于提升测试覆盖率
"""
import pytest
import sys
from docx import Document

sys.path.insert(0, 'D:/Projects/ThesisFormatCheck')


@pytest.mark.integration
class TestZoneDetectorBoundary:
    """ZoneDetector边界测试"""

    def test_empty_document(self):
        """边界测试：空文档"""
        from emba_checker.zone_detector import ZoneDetector
        
        doc = Document()
        detector = ZoneDetector(doc)
        
        assert detector.zone_map == {}

    def test_single_empty_paragraph(self):
        """边界测试：单个空段落"""
        from emba_checker.zone_detector import ZoneDetector
        
        doc = Document()
        doc.add_paragraph("")
        
        detector = ZoneDetector(doc)
        
        assert 0 in detector.zone_map

    def test_single_paragraph(self):
        """边界测试：单个段落"""
        from emba_checker.zone_detector import ZoneDetector
        
        doc = Document()
        doc.add_paragraph("测试内容")
        
        detector = ZoneDetector(doc)
        
        assert 0 in detector.zone_map
        assert detector.zone_map[0] == "cover"

    def test_many_paragraphs(self):
        """边界测试：多段落（100个）"""
        from emba_checker.zone_detector import ZoneDetector
        
        doc = Document()
        for i in range(100):
            doc.add_paragraph(f"段落 {i}")
        
        detector = ZoneDetector(doc)
        
        assert len(detector.zone_map) == 100

    def test_very_long_paragraph(self):
        """边界测试：超长段落"""
        from emba_checker.zone_detector import ZoneDetector
        
        doc = Document()
        long_text = "测试内容 " * 1000
        doc.add_paragraph(long_text)
        
        detector = ZoneDetector(doc)
        
        assert 0 in detector.zone_map

    def test_special_characters(self):
        """边界测试：特殊字符"""
        from emba_checker.zone_detector import ZoneDetector
        
        doc = Document()
        # 使用XML兼容的特殊字符，不使用控制字符
        doc.add_paragraph("特殊字符：中文、English、123、!@#$%^&*()")
        
        detector = ZoneDetector(doc)
        
        assert 0 in detector.zone_map

    def test_unicode_characters(self):
        """边界测试：Unicode字符"""
        from emba_checker.zone_detector import ZoneDetector
        
        doc = Document()
        doc.add_paragraph("中文\u4e00\u4e01 \U0001f600")
        
        detector = ZoneDetector(doc)
        
        assert 0 in detector.zone_map

    def test_whitespace_only(self):
        """边界测试：仅空白字符"""
        from emba_checker.zone_detector import ZoneDetector
        
        doc = Document()
        doc.add_paragraph("   \t\n   ")
        
        detector = ZoneDetector(doc)
        
        assert 0 in detector.zone_map

    def test_newline_only(self):
        """边界测试：仅换行符"""
        from emba_checker.zone_detector import ZoneDetector
        
        doc = Document()
        doc.add_paragraph("\n")
        
        detector = ZoneDetector(doc)
        
        assert 0 in detector.zone_map


@pytest.mark.integration
class TestZoneDetectorException:
    """ZoneDetector异常测试"""

    def test_empty_zone_map_attribute(self):
        """测试：zone_map属性"""
        from emba_checker.zone_detector import ZoneDetector
        
        doc = Document()
        doc.add_paragraph("测试")
        
        detector = ZoneDetector(doc)
        
        assert isinstance(detector.zone_map, dict)

    def test_detect_method(self):
        """测试：_detect方法"""
        from emba_checker.zone_detector import ZoneDetector
        
        doc = Document()
        doc.add_paragraph("测试")
        
        detector = ZoneDetector(doc)
        detector._detect()
        
        assert isinstance(detector.zone_map, dict)


@pytest.mark.integration
class TestZoneDetectorFullWorkflow:
    """ZoneDetector完整工作流测试"""

    def test_full_paper_flow(self):
        """测试：完整论文流程"""
        from emba_checker.zone_detector import ZoneDetector
        
        doc = Document()
        
        # 封面
        doc.add_paragraph("北京大学学位论文")
        
        # 版权声明
        doc.add_paragraph("版权声明")
        
        # 中文摘要
        doc.add_paragraph("摘要")
        doc.add_paragraph("这是中文摘要内容")
        
        # 英文摘要
        doc.add_paragraph("ABSTRACT")
        doc.add_paragraph("This is English abstract")
        
        # 目录
        doc.add_paragraph("目录")
        
        # 第一章
        doc.add_paragraph("第一章 导论")
        doc.add_paragraph("正文内容")
        
        # 第二章
        doc.add_paragraph("第二章 文献综述")
        doc.add_paragraph("文献内容")
        
        # 参考文献
        doc.add_paragraph("参考文献")
        
        # 致谢
        doc.add_paragraph("致谢")
        
        detector = ZoneDetector(doc)
        
        # 验证区域转换 - 段落索引可能因检测逻辑而变化
        assert detector.zone_map[0] == "cover"
        assert "copyright" in detector.zone_map.values()
        assert "abstract_cn" in detector.zone_map.values()
        assert "abstract_en" in detector.zone_map.values()
        assert "body_chapter" in detector.zone_map.values()
        assert "references" in detector.zone_map.values()
        assert "acknowledgment" in detector.zone_map.values()

    def test_minimal_paper(self):
        """测试：最小论文结构"""
        from emba_checker.zone_detector import ZoneDetector
        
        doc = Document()
        
        doc.add_paragraph("标题")
        doc.add_paragraph("正文")
        
        detector = ZoneDetector(doc)
        
        assert detector.zone_map[0] == "cover"
        assert detector.zone_map[1] == "cover"

    def test_only_chapters(self):
        """测试：只有章节的文档"""
        from emba_checker.zone_detector import ZoneDetector
        
        doc = Document()
        
        doc.add_paragraph("第一章")
        doc.add_paragraph("第二章")
        doc.add_paragraph("第三章")
        
        detector = ZoneDetector(doc)
        
        # 所有都应该是body_chapter
        for i in range(3):
            assert detector.zone_map[i] == "body_chapter"

    def test_mixed_content(self):
        """测试：混合内容"""
        from emba_checker.zone_detector import ZoneDetector
        
        doc = Document()
        
        # 封面
        doc.add_paragraph("封面标题")
        
        # 空段落
        doc.add_paragraph("")
        doc.add_paragraph("   ")
        
        # 正文
        doc.add_paragraph("第一章 绪论")
        doc.add_paragraph("正文段落")
        
        detector = ZoneDetector(doc)
        
        assert detector.zone_map[0] == "cover"
        assert detector.zone_map[1] == "cover"
        assert detector.zone_map[2] == "cover"
        assert detector.zone_map[3] == "body_chapter"
