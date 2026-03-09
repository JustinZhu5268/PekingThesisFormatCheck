# -*- coding: utf-8 -*-
import sys
sys.stdout.reconfigure(encoding='utf-8')

"""检查文档中是否还有未替换的引用"""
from docx import Document
import re

doc = Document('data/test_已修复_20260309_004116.docx')

print("=" * 60)
print("检查未替换的引用 [数字]")
print("=" * 60)

# 找到参考文献标题位置
ref_idx = None
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip().replace(' ', '')
    if text == "参考文献":
        ref_idx = i
        break

print(f"参考文献标题位置: {ref_idx}")

# 查找参考文献之前的所有未替换引用
unconverted = []
for i in range(ref_idx):
    para = doc.paragraphs[i]
    text = para.text
    # 查找 [数字] 格式的引用
    matches = re.findall(r'\[(\d+)\]', text)
    if matches:
        unconverted.append((i, matches, text[:80]))

print(f"\n发现 {len(unconverted)} 处未替换的引用:")
for para_idx, nums, text in unconverted[:20]:
    print(f"  段落 {para_idx}: {nums} - {text}...")

if len(unconverted) > 20:
    print(f"  ... 还有 {len(unconverted) - 20} 处")

print("\n" + "=" * 60)
print("检查正文是否还有上标格式")
print("=" * 60)

superscript_count = 0
for i in range(ref_idx):
    para = doc.paragraphs[i]
    for run in para.runs:
        if run.font.superscript is True:
            print(f"  段落 {i}, Run: 上标=True, 文本: {run.text[:50]}")
            superscript_count += 1

if superscript_count == 0:
    print("  未发现上标格式 ✓")
