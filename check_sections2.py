# -*- coding: utf-8 -*-
import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.oxml.ns import qn

doc = Document('data/test_已修复_20260309_030133.docx')

print('Total sections:', len(doc.sections))
for i, s in enumerate(doc.sections):
    print(f'\nSection {i}:')
    print(f'  top_margin: {s.top_margin.pt if s.top_margin else "N/A"}')
    print(f'  bottom_margin: {s.bottom_margin.pt if s.bottom_margin else "N/A"}')
    
    # 检查header
    header = s.header
    print(f'  header is_linked_to_previous: {header.is_linked_to_previous}')
    print(f'  header paragraphs: {len(header.paragraphs)}')
    if header.paragraphs:
        print(f'  header text: {header.paragraphs[0].text[:50] if header.paragraphs[0].text else "(empty)"}...')
    
    # 检查footer  
    footer = s.footer
    print(f'  footer is_linked_to_previous: {footer.is_linked_to_previous}')
    print(f'  footer paragraphs: {len(footer.paragraphs)}')
