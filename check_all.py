# -*- coding: utf-8 -*-
import sys
import io
if sys.platform == 'win32':
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8', errors='replace')

import glob
import os
from docx import Document

# Find the latest fixed file
data_dir = r'D:\Projects\ThesisFormatCheck\data'
files = glob.glob(os.path.join(data_dir, 'test_已修复_*.docx'))
files = [f for f in files if '备份' not in f]

if files:
    latest = max(files, key=os.path.getmtime)
    print(f"Latest file: {os.path.basename(latest)}")
    
    doc = Document(latest)
    
    print("\n=== 中文关键词 ===")
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if '关键词' in text and '：' in text:
            print(f"段落 {i}: {text}")
            if '；' in text:
                print("  [FAIL] 仍包含中文分号")
            if '，' in text:
                print("  [OK] 包含中文逗号")
    
    print("\n=== 英文关键词 ===")
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if 'KEY WORDS' in text or 'Keywords' in text:
            print(f"段落 {i}: {text}")
            if ';' in text:
                print("  [FAIL] 仍包含英文分号")
            if ',' in text:
                print("  [OK] 包含英文逗号")
    
    print("\n=== 参考文献示例 ===")
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if text.startswith('[') and ']' in text[:5]:
            print(f"段落 {i}: {text[:80]}")
            if i > 10:
                break
    
    print("\n=== 章节标题示例（检查2.4.2.格式）===")
    count = 0
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        # Check for headings like 2.4.2. or 1.2.3.
        import re
        if re.match(r'^\d+\.\d+\.\d+\.*', text):
            print(f"段落 {i}: {text[:60]}")
            count += 1
            if count >= 10:
                break
