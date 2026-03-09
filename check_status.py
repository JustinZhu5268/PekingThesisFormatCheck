# -*- coding: utf-8 -*-
import sys
import re
from docx import Document

# 检查修复后文件
doc = Document(r'D:\Projects\ThesisFormatCheck\data\test_已修复_20260307_202826.docx')

print('=== 检查未修复的图表编号格式 ===')
count_old = 0
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    matches = re.findall(r'[图表]\s*\d+-\d+', text)
    if matches:
        print(f'段落 {i}: {matches}')
        count_old += 1
print(f'旧格式(2-1)数量: {count_old}')

print()
print('=== 检查新格式图表编号 ===')
count_new = 0
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    matches = re.findall(r'[图表]\s*\d+\.\d+', text)
    if matches:
        count_new += 1
print(f'新格式(2.1)数量: {count_new}')
