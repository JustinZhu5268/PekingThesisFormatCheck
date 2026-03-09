# -*- coding: utf-8 -*-
"""验证修复后的文档 - 检查TOC是否正确显示"""
from docx import Document
from docx.oxml import OxmlElement

NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}

doc = Document('data/test_已修复_20260308_201607.docx')

print("=" * 70)
print("修复后文档 - TOC中的末尾章节")
print("=" * 70)

for i in range(175, 200):
    para = doc.paragraphs[i]
    style_name = para.style.name if para.style else 'None'
    
    if 'toc' in style_name.lower():
        text = para.text.strip()[:60]
        if text:
            print(f"段落 {i}: [{style_name}] {text}")

print("\n" + "=" * 70)
print("修复后文档 - Heading 1 末尾章节")
print("=" * 70)

for i in range(1000, len(doc.paragraphs)):
    para = doc.paragraphs[i]
    style_name = para.style.name if para.style else 'None'
    
    if 'Heading 1' in style_name:
        text = para.text.strip()[:60]
        
        # 检查是否有 numPr
        para_xml = para._element
        pPr = para_xml.find('.//w:pPr', NS)
        has_numPr = False
        if pPr is not None:
            numPr = pPr.find('.//w:numPr', NS)
            if numPr is not None:
                has_numPr = True
                numId = numPr.find('.//w:numId', NS)
                if numId is not None:
                    val = numId.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val')
                    print(f"段落 {i}: [{style_name}] {text} - numId={val}")
                else:
                    print(f"段落 {i}: [{style_name}] {text} - numPr存在但无numId")
            else:
                print(f"段落 {i}: [{style_name}] {text}")
        else:
            print(f"段落 {i}: [{style_name}] {text}")
