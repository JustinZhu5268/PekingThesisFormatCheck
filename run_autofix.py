# -*- coding: utf-8 -*-
"""
EMBA 论文格式审查工具 - 自动修复 CLI 入口
"""

import sys
import os
import glob
import io
import argparse

# 设置UTF-8输出
if sys.platform == 'win32':
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8', errors='replace')

# 添加项目根目录到路径
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, BASE_DIR)


def parse_args():
    """解析命令行参数"""
    parser = argparse.ArgumentParser(
        description='EMBA 论文格式自动修复工具',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
使用示例:
  python run_autofix.py input.docx                    # 修复文档
  python run_autofix.py input.docx -o ./output        # 指定输出目录
  python run_autofix.py input.docx --no-backup       # 不备份原文件
  python run_autofix.py input.docx -v                # 显示详细日志
        """
    )
    
    # 位置参数
    parser.add_argument(
        'input_file',
        nargs='?',
        help='输入的 Word 文档路径 (.docx)'
    )
    
    # 可选参数
    parser.add_argument(
        '-o', '--output',
        default=None,
        help='输出目录 (默认: 当前目录)'
    )
    
    parser.add_argument(
        '--no-backup',
        action='store_true',
        help='不备份原文件'
    )
    
    parser.add_argument(
        '-v', '--verbose',
        action='store_true',
        help='显示详细日志'
    )
    
    parser.add_argument(
        '-r', '--rules',
        default=None,
        help='规则文件路径 (默认: rules_registry.json)'
    )
    
    parser.add_argument(
        '-c', '--check',
        action='store_true',
        help='只检查格式问题，不修复（验证模式）'
    )

    parser.add_argument(
        '--diff',
        action='store_true',
        help='对比原始文档和修复后文档的格式差异（需要提供两个文件）'
    )

    parser.add_argument(
        '--inspect',
        action='store_true',
        help='检查文档的实际格式（封面、标题、关键词等）'
    )

    parser.add_argument(
        '--verify',
        action='store_true',
        help='验证修复结果：对比原始文档和修复后文档的内容差异'
    )

    parser.add_argument(
        '--original',
        default=None,
        help='原始文档路径（用于--diff或--verify模式）'
    )

    return parser.parse_args()


def run_check_mode(input_file: str, rules_path: str, verbose: bool = False):
    """检查模式：只检查格式问题，不修复"""
    import json
    from docx import Document
    from emba_checker.zone_detector import ZoneDetector
    from emba_checker.docx_engine import DocxEngine
    
    print("\n开始格式检查...\n")
    
    # 加载文档
    doc = Document(input_file)
    
    # 区域检测
    zone_detector = ZoneDetector(doc)
    zone_map = zone_detector.zone_map
    
    # 加载规则
    with open(rules_path, 'r', encoding='utf-8') as f:
        rules = json.load(f)
    
    # 只选择自动执行的规则
    enabled_rules = [r for r in rules if r.get("enabled", True)]
    
    print(f"加载了 {len(enabled_rules)} 条规则")
    print(f"文档共有 {len(doc.paragraphs)} 个段落\n")
    
    # 运行检查
    engine = DocxEngine(doc, enabled_rules, zone_map)
    issues = engine.run_all()
    
    # 统计结果
    error_count = sum(1 for i in issues if getattr(i, 'severity', '') == 'error')
    warning_count = sum(1 for i in issues if getattr(i, 'severity', '') == 'warning')
    info_count = sum(1 for i in issues if getattr(i, 'severity', '') == 'info')
    
    print("=" * 60)
    print("格式检查结果")
    print("=" * 60)
    print(f"发现问题总数: {len(issues)}")
    print(f"  错误: {error_count}")
    print(f"  警告: {warning_count}")
    print(f"  信息: {info_count}")
    print("=" * 60)
    
    # 显示所有问题
    if issues:
        print("\n详细问题列表:")
        print("-" * 60)
        for i, issue in enumerate(issues, 1):
            rule_id = getattr(issue, 'rule_id', '?')
            severity = getattr(issue, 'severity', 'info')
            message = getattr(issue, 'message', '')
            # 简化消息显示
            message = message.split('\n')[0][:80] if message else ""
            print(f"{i:3}. [{severity}] {rule_id}: {message}")
    else:
        print("\n未发现问题！")
    
    return 0 if not issues else 1


def run_diff_mode(input_file: str, original_file: str = None):
    """对比模式：对比原始文档和修复后文档的格式差异"""
    from docx import Document
    from emba_checker import utils

    print("\n开始格式对比...\n")

    # 如果没有指定原始文件，尝试自动查找
    if not original_file:
        # 尝试查找data目录下的原始文件
        data_dir = os.path.dirname(input_file) if input_file else 'data'
        if not data_dir:
            data_dir = 'data'

        # 查找可能的原始文件（不含"已修复"或"_fixed"字样）
        import glob as g
        all_files = g.glob(os.path.join(data_dir, '*.docx'))

        for f in all_files:
            basename = os.path.basename(f)
            if '已修复' not in basename and '_fixed' not in basename and basename != os.path.basename(input_file):
                original_file = f
                break

    if not original_file or not os.path.exists(original_file):
        print(f"错误: 原始文件不存在: {original_file}", file=sys.stderr)
        return 1

    print(f"原始文件: {original_file}")
    print(f"修复后文件: {input_file}\n")

    doc_orig = Document(original_file)
    doc_fixed = Document(input_file)

    # 检查正文段落对齐方式
    print("=" * 60)
    print("正文段落对齐方式对比")
    print("=" * 60)
    for i in range(20, 40):
        text = doc_orig.paragraphs[i].text.strip()
        if len(text) > 50:
            align_orig = utils.get_paragraph_alignment(doc_orig.paragraphs[i])
            align_fixed = utils.get_paragraph_alignment(doc_fixed.paragraphs[i])
            indent_orig = utils.get_first_line_indent(doc_orig.paragraphs[i])
            indent_fixed = utils.get_first_line_indent(doc_fixed.paragraphs[i])

            print(f"段落 {i}: {text[:40]}...")
            print(f"  原始对齐: {align_orig}, 首行缩进: {indent_orig}")
            print(f"  修复后对齐: {align_fixed}, 首行缩进: {indent_fixed}")

            if align_orig != align_fixed:
                print(f"  ❌ 对齐方式被改变了!")
            if indent_orig != indent_fixed:
                print(f"  ❌ 首行缩进被改变了!")
            print()
            break

    # 检查关键词
    print("=" * 60)
    print("关键词对比")
    print("=" * 60)
    for i in range(len(doc_orig.paragraphs)):
        text = doc_orig.paragraphs[i].text.strip()
        if "关键词" in text and "：" in text:
            print(f"原始: {text}")
            print(f"修复后: {doc_fixed.paragraphs[i].text.strip()}")
            print()
            break

    # 检查摘要标题格式
    print("=" * 60)
    print("摘要标题格式对比")
    print("=" * 60)
    for i in range(len(doc_orig.paragraphs)):
        text = doc_orig.paragraphs[i].text.strip()
        if text == "摘要":
            print(f"段落 {i}: {text}")
            align_orig = utils.get_paragraph_alignment(doc_orig.paragraphs[i])
            align_fixed = utils.get_paragraph_alignment(doc_fixed.paragraphs[i])
            for run in doc_orig.paragraphs[i].runs:
                if run.text.strip():
                    size_orig = utils.get_font_size(run)
                    break
            for run in doc_fixed.paragraphs[i].runs:
                if run.text.strip():
                    size_fixed = utils.get_font_size(run)
                    break
            print(f"  原始对齐: {align_orig}, 字号: {size_orig}")
            print(f"  修复后对齐: {align_fixed}, 字号: {size_fixed}")
            print()
            break

    return 0


def run_verify_mode(fixed_file: str, original_file: str = None):
    """验证模式：对比原始文档和修复后文档的内容差异"""
    import re
    from docx import Document
    
    print("\n" + "=" * 60)
    print("验证修复结果 - 内容对比")
    print("=" * 60)
    
    data_dir = os.path.join(BASE_DIR, 'data')
    
    # 自动查找最新的修复文件和对应备份
    if not original_file:
        fixed_files = [f for f in glob.glob(os.path.join(data_dir, 'test_已修复_*.docx')) 
                      if '原文件' not in f and '备份' not in f]
        if fixed_files:
            # 按修改时间排序，找最新的
            fixed_file = max(fixed_files, key=os.path.getmtime)
            # 对应的备份文件
            timestamp = os.path.basename(fixed_file).replace('test_已修复_', '').replace('.docx', '')
            original_file = os.path.join(data_dir, f'test_已修复_{timestamp}_原文件备份.docx')
            print(f"修复文件: {os.path.basename(fixed_file)}")
            print(f"备份文件: {os.path.basename(original_file)}")
    
    if not original_file or not os.path.exists(original_file):
        print("错误: 未找到原始文档，请使用 --original 指定")
        return 1
    
    print(f"\n比较文档:")
    print(f"  原始: {os.path.basename(original_file)}")
    print(f"  修复: {os.path.basename(fixed_file)}")
    print()
    
    doc1 = Document(original_file)
    doc2 = Document(fixed_file)
    
    paragraphs1 = [p.text.strip() for p in doc1.paragraphs]
    paragraphs2 = [p.text.strip() for p in doc2.paragraphs]
    
    print(f"原始文档段落数: {len(paragraphs1)}")
    print(f"修复文档段落数: {len(paragraphs2)}")
    print("=" * 60)
    
    # 标准化函数
    def normalize_for_comparison(text):
        """标准化文本用于比较：排除格式差异"""
        # 图表编号格式：图2-1 -> 图2.1
        text = re.sub(r'([图表])\s*(\d+)\s*-\s*(\d+)', r'\1\2.\3', text)
        # 中文标点转英文
        text = text.replace('；', ',').replace('：', ':')
        # 去除标题后的小数点：1.1. 标题 -> 1.1 标题
        text = re.sub(r'^(\d+\.\d+)\.\s+', r'\1 ', text)
        # 去除三级标题后的小数点
        text = re.sub(r'^(\d+\.\d+\.\d+)\.\s+', r'\1 ', text)
        # 去掉多余空格
        text = re.sub(r'\s+', '', text)
        return text
    
    # 找出差异
    differences = []
    format_changes = []
    
    max_len = max(len(paragraphs1), len(paragraphs2))
    
    for i in range(max_len):
        p1 = paragraphs1[i] if i < len(paragraphs1) else "[不存在]"
        p2 = paragraphs2[i] if i < len(paragraphs2) else "[不存在]"
        
        if p1 != p2:
            norm1 = normalize_for_comparison(p1)
            norm2 = normalize_for_comparison(p2)
            
            if norm1 == norm2:
                format_changes.append({'index': i, 'original': p1, 'fixed': p2})
            else:
                differences.append({
                    'index': i,
                    'original': p1,
                    'fixed': p2,
                    'original_len': len(p1),
                    'fixed_len': len(p2)
                })
    
    print(f"\n发现 {len(differences)} 处文字差异")
    print(f"发现 {len(format_changes)} 处格式变化（已排除）\n")
    
    # 显示前30个差异
    for i, diff in enumerate(differences[:30]):
        print(f"{i+1}. 段落 {diff['index']}:")
        print(f"   原始: {diff['original'][:60]}...")
        print(f"   修复: {diff['fixed'][:60]}...")
        print()
    
    if len(differences) > 30:
        print(f"... 还有 {len(differences) - 30} 处差异")
    
    # 统计格式变化类型
    fig_num_changes = sum(1 for fc in format_changes 
                         if re.search(r'[图表]\d+-\d+', fc['original']) or re.search(r'[图表]\d+\.\d+', fc['fixed']))
    
    print("\n" + "=" * 60)
    print("格式变化统计:")
    print("=" * 60)
    print(f"图表编号格式变化: {fig_num_changes} 处")
    print(f"其他格式变化: {len(format_changes) - fig_num_changes} 处")
    print(f"文字差异: {len(differences)} 处")
    
    # 检查是否是预期的变化
    expected_keyword_fix = False
    expected_ref_fix = False
    
    for diff in differences:
        orig = diff['original']
        fixed = diff['fixed']
        idx = diff['index']
        
        # 关键词分隔符变化
        if '关键词' in orig and '；' in orig and '，' in fixed and '；' not in fixed:
            expected_keyword_fix = True
        
        # 参考文献区域变化
        if idx > 1000:
            expected_ref_fix = True
    
    # 判断结果
    print("\n" + "=" * 60)
    print("结论:")
    print("=" * 60)

    # 只有关键词或参考文献区域的变化，且变化数量合理
    has_expected_changes = expected_keyword_fix or expected_ref_fix
    # 允许参考文献区域最多50处变化（每个条目一处），关键词最多1处
    reasonable_ref_changes = expected_ref_fix and len(differences) <= 50
    reasonable_keyword_changes = expected_keyword_fix and len(differences) <= 1
    only_expected = has_expected_changes and (reasonable_ref_changes or reasonable_keyword_changes)
    
    if only_expected:
        print("✓ 内容变化仅为预期的格式修复")
        print("  (关键词分隔符 + 参考文献格式)")
        return 0
    elif len(differences) == 0:
        print("✓ 内容变化仅限于格式，没有实质性文字变化")
        return 0
    else:
        print(f"✗ 发现 {len(differences)} 处非预期的实质性变化！")
        print("  需要检查修复逻辑")
        return 1


def run_inspect_mode(input_file: str):
    """检查模式：检查文档的实际格式"""
    from docx import Document
    from emba_checker import utils

    print("\n开始检查文档格式...\n")

    doc = Document(input_file)

    # 检查封面标题
    print("=" * 60)
    print("封面标题")
    print("=" * 60)
    for i, para in enumerate(doc.paragraphs[:10]):
        if para.text.strip():
            print(f"段落 {i}: {para.text[:50]}...")
            # 检查字体
            for run in para.runs:
                if run.text.strip():
                    font_name = utils.get_font_name(run)
                    font_size = utils.get_font_size(run)
                    print(f"  字体: {font_name}")
                    print(f"  字号: {font_size}")
                    break
            # 检查对齐
            align = utils.get_paragraph_alignment(para)
            print(f"  对齐: {align}")
            print()

    # 检查中英文摘要标题
    print("=" * 60)
    print("摘要标题")
    print("=" * 60)
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if text in ["摘要", "ABSTRACT", "关键词", "Keywords"]:
            print(f"段落 {i}: {text}")
            align = utils.get_paragraph_alignment(para)
            print(f"  对齐: {align}")
            for run in para.runs:
                if run.text.strip():
                    font_name = utils.get_font_name(run)
                    font_size = utils.get_font_size(run)
                    bold = utils.get_font_bold(run)
                    print(f"  字体: {font_name}, 字号: {font_size}, 加粗: {bold}")
                    break
            print()

    # 检查关键词行
    print("=" * 60)
    print("关键词段落")
    print("=" * 60)
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if "关键词" in text and "：" in text:
            print(f"段落 {i}: {text[:60]}...")
            # 检查是否有分号
            if "；" in text:
                print("  ❌ 仍包含中文分号")
            if ";" in text:
                print("  ❌ 仍包含英文分号")
            if "，" in text:
                print("  ✓ 包含中文逗号")
            print()

    # 检查章节标题
    print("=" * 60)
    print("章节标题")
    print("=" * 60)
    chapter_count = 0
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if text.startswith("一、") or text.startswith("二、") or text.startswith("三、") or \
           text.startswith("四、") or text.startswith("五、") or text.startswith("六、") or \
           text.startswith("七、") or text.startswith("八、") or text.startswith("九、") or \
           text.startswith("十、"):
            print(f"段落 {i}: {text[:40]}...")
            align = utils.get_paragraph_alignment(para)
            print(f"  对齐: {align}")
            for run in para.runs:
                if run.text.strip():
                    font_name = utils.get_font_name(run)
                    font_size = utils.get_font_size(run)
                    bold = utils.get_font_bold(run)
                    print(f"  字体: {font_name}, 字号: {font_size}, 加粗: {bold}")
                    break
            print()
            chapter_count += 1
            if chapter_count >= 5:
                break

    # 检查正文段落格式
    print("=" * 60)
    print("正文段落示例")
    print("=" * 60)
    for i, para in enumerate(doc.paragraphs[20:40]):
        text = para.text.strip()
        if len(text) > 50:
            print(f"段落 {i+20}: {text[:50]}...")
            align = utils.get_paragraph_alignment(para)
            indent = utils.get_first_line_indent(para)
            line_spacing = utils.get_line_spacing(para)
            print(f"  对齐: {align}")
            print(f"  首行缩进: {indent}")
            print(f"  行距: {line_spacing}")
            print()
            break

    return 0


def run_fix_mode(input_file: str, rules_path: str, args):
    """修复模式：自动修复格式问题"""
    import re
    from docx import Document
    from emba_checker.docx_autofixer import auto_fix

    print("\n" + "=" * 60)
    print("开始自动修复...")
    print("=" * 60)
    
    # ========== 修复前检查 ==========
    print("\n【修复前格式检查】")
    doc = Document(input_file)
    
    # 1. 检查节标题小数点格式
    heading_issues = []
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if re.match(r'^\d+\.\d+\.\s+', text):
            heading_issues.append((i, text[:50]))
    
    # 2. 检查图表编号格式
    figure_issues = []
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        matches = re.findall(r'[图表]\s*\d+-\d+', text)
        if matches:
            figure_issues.append((i, matches))
    
    # 3. 检查参考文献
    ref_issues = []
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if re.match(r'^\s*\[\d+\]', text):
            ref_issues.append((i, text[:50]))
    
    print(f"  节标题小数点问题: {len(heading_issues)} 处")
    print(f"  图表编号格式问题: {len(figure_issues)} 处")
    print(f"  参考文献序号问题: {len(ref_issues)} 处")
    print("-" * 60)
    
    # 执行修复
    output_path, report = auto_fix(
        input_file,
        rules_path,
        output_path=args.output,
        backup=not args.no_backup,
        verbose=args.verbose
    )
    
    # ========== 修复后检查 ==========
    print("\n【修复后验证】")
    doc_fixed = Document(output_path)
    
    # 1. 检查节标题
    heading_fixed = 0
    for i, para in enumerate(doc_fixed.paragraphs):
        text = para.text.strip()
        if re.match(r'^\d+\.\d+\.\s+', text):
            heading_fixed += 1
    
    # 2. 检查图表编号
    figure_fixed = 0
    for i, para in enumerate(doc_fixed.paragraphs):
        text = para.text.strip()
        matches = re.findall(r'[图表]\s*\d+-\d+', text)
        if matches:
            figure_fixed += 1
    
    # 3. 检查参考文献
    ref_fixed = 0
    for i, para in enumerate(doc_fixed.paragraphs):
        text = para.text.strip()
        if re.match(r'^\s*\[\d+\]', text):
            ref_fixed += 1
    
    # 输出摘要
    summary = report.get_summary()

    print("\n" + "=" * 60)
    print("修复完成")
    print("=" * 60)
    print(f"输出文件: {output_path}")
    print(f"总修复项: {summary['total_fixes']}")
    print(f"  成功: {summary['successful_fixes']}")
    print(f"  失败: {summary['failed_fixes']}")
    print("-" * 60)
    print("\n格式修复对比:")
    print(f"  节标题小数点: {len(heading_issues)} -> {heading_fixed}")
    print(f"  图表编号: {len(figure_issues)} -> {figure_fixed}")
    print(f"  参考文献序号: {len(ref_issues)} -> {ref_fixed}")
    print("=" * 60)

    # 生成修复报告文件
    report_path = output_path.replace(".docx", "_修复报告.txt")
    with open(report_path, 'w', encoding='utf-8') as f:
        f.write(report.to_string())

    print(f"\n修复报告: {report_path}")

    return 0


def main():
    """主入口"""
    args = parse_args()
    
    # 自动查找docx文件
    docx_files = glob.glob('data/*.docx')
    
    # 显示帮助信息
    if args.input_file is None and not docx_files:
        print("""
EMBA 论文格式自动修复工具

使用方法:
  python run_autofix.py input.docx                    # 修复文档
  python run_autofix.py input.docx --check            # 只检查不修复
  python run_autofix.py input.docx --inspect          # 检查文档实际格式
  python run_autofix.py input.docx --diff             # 对比格式差异
  python run_autofix.py input.docx --diff --original original.docx  # 指定原始文件对比
  python run_autofix.py input.docx -o ./output        # 指定输出目录
  python run_autofix.py input.docx --no-backup        # 不备份原文件
  python run_autofix.py input.docx -v                  # 显示详细日志
        """)
        return 0
    
    # 使用输入参数或自动查找
    if args.input_file:
        input_file = args.input_file
    elif docx_files:
        input_file = docx_files[0]
        print(f"自动选择文件: {input_file}")
    else:
        print("错误: 未找到输入文件", file=sys.stderr)
        return 1
    
    # 验证输入文件
    if not os.path.exists(input_file):
        print(f"错误: 输入文件不存在: {input_file}", file=sys.stderr)
        return 1
    
    # 验证文件扩展名
    if not input_file.lower().endswith('.docx'):
        print(f"错误: 输入文件必须是 .docx 格式", file=sys.stderr)
        return 1
    
    # 获取规则文件路径
    rules_path = args.rules if args.rules else os.path.join(BASE_DIR, 'emba_checker', 'rules_registry.json')
    if not os.path.exists(rules_path):
        print(f"错误: 规则文件不存在: {rules_path}", file=sys.stderr)
        return 1
    
    # 打印配置信息
    print("=" * 60)
    print("EMBA 论文格式自动修复工具")
    print("=" * 60)
    print(f"论文文件: {input_file}")
    print(f"规则文件: {rules_path}")
    print(f"备份原文件: {'否' if args.no_backup else '是'}")
    if args.output:
        print(f"输出目录: {args.output}")

    # 确定运行模式
    if args.diff:
        mode = "对比模式（对比原始和修复文档格式差异）"
    elif args.verify:
        mode = "验证模式（对比原始和修复文档内容）"
    elif args.inspect:
        mode = "检查模式（检查文档实际格式）"
    elif args.check:
        mode = "检查模式（只检查不修复）"
    else:
        mode = "修复模式"
    print(f"运行模式: {mode}")
    print("-" * 60)

    try:
        # 对比模式
        if args.diff:
            return run_diff_mode(input_file, args.original)
        # 验证模式
        elif args.verify:
            return run_verify_mode(input_file, args.original)
        # 检查模式（检查文档实际格式）
        elif args.inspect:
            return run_inspect_mode(input_file)
        # 检查模式（验证格式问题）
        elif args.check:
            return run_check_mode(input_file, rules_path, args.verbose)
        else:
            # 修复模式 - 调用自动修复
            return run_fix_mode(input_file, rules_path, args)
        
    except KeyboardInterrupt:
        print("\n\n操作已取消", file=sys.stderr)
        return 130
        
    except Exception as e:
        print(f"\n错误: {str(e)}", file=sys.stderr)
        if args.verbose:
            import traceback
            traceback.print_exc()
        return 1


if __name__ == "__main__":
    sys.exit(main())
