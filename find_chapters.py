# -*- coding: utf-8 -*-
"""查找真正包含"第八章 参考文献"等文本的段落"""
from docx import Document

doc = Document('data/test.docx')

print("=" * 70)
print("查找包含'第八章'/'第九章'/'第十章'的所有段落")
print("=" * 70)

for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    style_name = para.style.name if para.style else 'None'
    
    if '第八章' in text or '第九章' in text or '第十章' in text:
        print(f"\n段落 {i}:")
        print(f"  样式: {style_name}")
        print(f"  文本: {text[:80]}")
        
        # 打印runs
        for j, run in enumerate(para.runs):
            print(f"    Run {j}: '{run.text}'")

print("\n" + "=" * 70)
print("查找末尾的'参考文献'/'附录'/'致谢'段落(Heading 1)")
print("=" * 70)

for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    style_name = para.style.name if para.style else 'None'
    
    if style_name == 'Heading 1' and ('参考文献' in text or '附录' in text or '致谢' in text):
        print(f"\n段落 {i}:")
        print(f"  样式: {style_name}")
        print(f"  文本: {text[:80]}")
