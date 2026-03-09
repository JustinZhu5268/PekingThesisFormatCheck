# -*- coding: utf-8 -*-
"""研究Word Heading样式的XML结构"""
import codecs
from docx import Document
from docx.oxml import OxmlElement

NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}

output = []

doc = Document('data/test.docx')

output.append("=" * 60)
output.append("查找Heading 1段落")
output.append("=" * 60)

# 检查前面章节的Heading 1
for i in range(min(250, len(doc.paragraphs))):
    para = doc.paragraphs[i]
    style_name = para.style.name if para.style else 'None'
    if 'Heading 1' in style_name:
        text = para.text.strip()[:40]
        para_xml = para._element
        numPr = para_xml.find('.//w:numPr', NS)
        numPr_str = "有numPr" if numPr is not None else "无numPr"
        output.append(f"段落 {i}: [{style_name}] {text} - {numPr_str}")

output.append("\n" + "=" * 60)
output.append("查找末尾章节")
output.append("=" * 60)

# 检查末尾章节
for i in range(1000, len(doc.paragraphs)):
    para = doc.paragraphs[i]
    style_name = para.style.name if para.style else 'None'
    text = para.text.strip()[:40]
    if '参考文献' in text or '附录' in text or '致谢' in text:
        para_xml = para._element
        numPr = para_xml.find('.//w:numPr', NS)
        numPr_str = "有numPr" if numPr is not None else "无numPr"
        output.append(f"段落 {i}: [{style_name}] {text} - {numPr_str}")

# 保存到文件
with open('study_output.txt', 'w', encoding='utf-8') as f:
    f.write('\n'.join(output))

print("已保存到 study_output.txt")
