# -*- coding: utf-8 -*-
import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document

doc = Document('data/test_已修复_20260309_094827.docx')

print('Total sections:', len(doc.sections))
for i, s in enumerate(doc.sections):
    header = s.header
    header_text = header.paragraphs[0].text[:40] if header.paragraphs and header.paragraphs[0].text else '(empty)'
    print(f'Section {i}: header="{header_text}"')
