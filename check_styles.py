# -*- coding: utf-8 -*-
"""检查文档中的标题样式"""
import os
import glob
from docx import Document

# 找到最新的修复文件
docx_files = glob.glob('data/test_已修复_*.docx')
# 排除备份文件
docx_files = [f for f in docx_files if '_原文件备份' not in f]
latest = sorted(docx_files)[-1]
print(f"检查文件: {latest}")

doc = Document(latest)

# 查找"第一章"段落及其样式
print("\n查找'第一章'段落:")
for i, para in enumerate(doc.paragraphs):
    if para.text.strip().startswith('第一章'):
        print(f"  段落 {i}: '{para.text[:30]}...'")
        print(f"    样式: {para.style.name}")
        break

# 查找所有使用的样式
print("\n文档中使用的样式:")
styles = set()
for para in doc.paragraphs:
    if para.style.name:
        styles.add(para.style.name)
        
for s in sorted(styles):
    if '标题' in s or 'Heading' in s or 'head' in s.lower():
        print(f"  {s}")
