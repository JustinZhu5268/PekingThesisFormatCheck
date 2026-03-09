# -*- coding: utf-8 -*-
import sys
sys.stdout.reconfigure(encoding='utf-8')

"""验证修复后的文档"""
from docx import Document
from docx.shared import Pt

doc = Document('data/test_已修复_20260309_002901.docx')

print("=" * 60)
print("检查正文中的引用格式（是否还有上标）")
print("=" * 60)

# 查找正文中的引用，检查上标属性
for i in [272, 279, 284, 287, 290]:
    para = doc.paragraphs[i]
    for run in para.runs:
        if '（' in run.text and '）' in run.text:
            print(f"段落 {i}, Run: 上标={run.font.superscript}, 文本: {run.text[:40]}")

print("\n" + "=" * 60)
print("检查参考文献段落格式")
print("=" * 60)

# 找到参考文献标题
ref_idx = None
for i, para in enumerate(doc.paragraphs):
    if para.text.strip() == "参考文献":
        ref_idx = i
        break

if ref_idx:
    print(f"参考文献标题在段落 {ref_idx}")
    print("\n前3条参考文献格式:")
    for i in range(ref_idx + 1, min(ref_idx + 4, len(doc.paragraphs))):
        para = doc.paragraphs[i]
        if not para.text.strip():
            continue
        
        # 检查格式
        pf = para.paragraph_format
        run = para.runs[0] if para.runs else None
        
        # 转换字号
        font_size_pt = None
        if run and run.font.size:
            font_size_pt = run.font.size.pt
            
        print(f"\n段落 {i}: {para.text[:50]}...")
        print(f"  字体: {run.font.name if run else 'N/A'}")
        print(f"  字号: {font_size_pt} pt" if font_size_pt else "  字号: N/A")
        
        # 转换行距
        line_spacing_pt = None
        if pf.line_spacing:
            line_spacing_pt = pf.line_spacing.pt if hasattr(pf.line_spacing, 'pt') else pf.line_spacing / 12700
        print(f"  行距: {line_spacing_pt} pt" if line_spacing_pt else f"  行距: {pf.line_spacing}")
        print(f"  行距规则: {pf.line_spacing_rule}")
        
        # 转换段前段后
        space_before_pt = pf.space_before.pt if pf.space_before and hasattr(pf.space_before, 'pt') else 0
        space_after_pt = pf.space_after.pt if pf.space_after and hasattr(pf.space_after, 'pt') else 0
        print(f"  段前: {space_before_pt} pt")
        print(f"  段后: {space_after_pt} pt")
        
        # 转换缩进
        left_indent_pt = pf.left_indent.pt if pf.left_indent and hasattr(pf.left_indent, 'pt') else 0
        first_line_pt = pf.first_line_indent.pt if pf.first_line_indent and hasattr(pf.first_line_indent, 'pt') else 0
        print(f"  左缩进: {left_indent_pt} pt")
        print(f"  首行缩进: {first_line_pt} pt")
